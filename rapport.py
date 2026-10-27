# -*- coding: utf-8 -*-
"""
Rapport automatique de StatExcel : hypothèses, exécution des tests, rédaction et export Word.

Le rapport est d'abord construit comme une liste de « blocs » (titres, paragraphes, tableaux,
figures), affichée telle quelle dans l'application puis écrite dans un fichier .docx.

Auteur : BOKOUMBO Tibetobe Ambroise
"""
import datetime
import io

import numpy as np
import pandas as pd
from scipy import stats

import figures as fg
import stats_utils as su

MOIS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août",
        "septembre", "octobre", "novembre", "décembre"]

TYPES_TEST = {
    "groupes": "Comparaison de groupes (t, Mann-Whitney, ANOVA, Kruskal-Wallis)",
    "khi2": "Association entre deux variables qualitatives (khi-deux / Fisher)",
    "correlation": "Corrélation entre deux variables numériques",
    "apparie": "Comparaison de deux mesures appariées (t apparié / Wilcoxon)",
    "normalite": "Normalité d'une variable",
    "regression": "Régression linéaire multiple (OLS)",
    "fiabilite": "Fiabilité d'une échelle (alpha de Cronbach)",
}


# ---------------------------------------------------------------- outils de rédaction
def n2(x, d=2):
    if x is None or pd.isna(x):
        return "—"
    return f"{x:.{d}f}".replace(".", ",")


def _vf(s):
    return str(s).replace(".", ",")


def date_fr(d=None):
    d = d or datetime.date.today()
    return f"{d.day} {MOIS[d.month - 1]} {d.year}"


def H(niveau, texte):
    return {"type": f"h{niveau}", "texte": texte}


def P(texte, gras=None):
    return {"type": "p", "texte": texte, "gras": gras}


def L(items):
    return {"type": "liste", "items": list(items)}


def T(df, legende):
    return {"type": "table", "df": df, "legende": legende}


def F(png, legende):
    return {"type": "fig", "png": png, "legende": legende}


def _paragraphes(texte):
    return [P(t.strip()) for t in (texte or "").split("\n") if t.strip()]


def _conclusion_p(p, alpha):
    a = n2(alpha, 2)
    tp = su.texte_p(p)
    if p < alpha:
        return (f"Au seuil α = {a}, le résultat est statistiquement significatif ({tp}) : "
                f"l'hypothèse nulle H0 est rejetée en faveur de H1.")
    return (f"Au seuil α = {a}, le résultat n'est pas statistiquement significatif ({tp}) : "
            f"H0 ne peut pas être rejetée. Cela ne prouve pas que H0 est vraie ; les données "
            f"ne permettent simplement pas de conclure à un effet.")


def _decision(p, alpha):
    if p is None or pd.isna(p):
        return "Non calculable"
    return "H0 rejetée" if p < alpha else "H0 non rejetée"


# ---------------------------------------------------------------- hypothèses par défaut
def titre_hypothese(h):
    p = h["params"]
    t = h["type"]
    if t == "groupes":
        return f"{p['y']} selon {p['g']}"
    if t == "khi2":
        return f"{p['a']} et {p['b']}"
    if t == "correlation":
        return f"{p['x']} et {p['y']}"
    if t == "apparie":
        return f"{p['a']} et {p['b']} (mesures appariées)"
    if t == "normalite":
        return f"Normalité de {p['v']}"
    if t == "regression":
        return f"Déterminants de {p['y']}"
    return "Fiabilité de l'échelle (" + ", ".join(p["items"][:4]) + ("…" if len(p["items"]) > 4 else "") + ")"


def hypotheses_par_defaut(type_h, p, nb_groupes=2):
    """Propose (H0, H1) ; l'utilisateur peut les modifier."""
    if type_h == "groupes":
        if nb_groupes <= 2:
            return (f"{p['y']} est identique dans les deux groupes de {p['g']}.",
                    f"{p['y']} diffère entre les deux groupes de {p['g']}.")
        return (f"{p['y']} est identique dans tous les groupes de {p['g']}.",
                f"Au moins un groupe de {p['g']} diffère des autres pour {p['y']}.")
    if type_h == "khi2":
        return (f"Les variables {p['a']} et {p['b']} sont indépendantes.",
                f"Les variables {p['a']} et {p['b']} sont associées (dépendantes).")
    if type_h == "correlation":
        return (f"Il n'existe pas de corrélation entre {p['x']} et {p['y']} (coefficient = 0).",
                f"Il existe une corrélation entre {p['x']} et {p['y']} (coefficient ≠ 0).")
    if type_h == "apparie":
        return (f"Les valeurs de {p['a']} et de {p['b']} sont identiques en moyenne (différence nulle).",
                f"Les valeurs de {p['a']} et de {p['b']} diffèrent.")
    if type_h == "normalite":
        return (f"{p['v']} suit une loi normale.", f"{p['v']} ne suit pas une loi normale.")
    if type_h == "regression":
        return (f"Aucune variable explicative n'a d'effet linéaire sur {p['y']} (tous les coefficients sont nuls).",
                f"Au moins une variable explicative a un effet linéaire sur {p['y']}.")
    return ("L'échelle n'est pas suffisamment fiable (alpha de Cronbach < 0,70).",
            "L'échelle est fiable (alpha de Cronbach ≥ 0,70).")


# ---------------------------------------------------------------- tests (un par type)
def _t_groupes(df, p, alpha, avec_fig):
    y, g = p["y"], p["g"]
    groupes = {k: v for k, v in su.groupes_numeriques(df, y, g).items() if len(v) > 0}
    if len(groupes) < 2:
        raise ValueError("il faut au moins deux groupes contenant des observations")
    if any(len(v) < 2 for v in groupes.values()):
        raise ValueError("chaque groupe doit contenir au moins 2 observations")
    noms, k = list(groupes), len(groupes)
    effectifs = ", ".join(f"« {n} » : n = {len(groupes[n])}" for n in noms)
    norm = pd.DataFrame([{"Groupe": n, "N": len(v),
                          "Shapiro-Wilk (p)": stats.shapiro(v)[1] if 3 <= len(v) <= 5000 and np.ptp(v) > 0 else np.nan}
                         for n, v in groupes.items()]).set_index("Groupe")
    blocs = [P(f"La variable quantitative « {y} » est comparée selon les {k} modalités de « {g} » ({effectifs})."),
             T(su.descriptifs_groupes(groupes), f"Statistiques descriptives de {y} selon {g}")]
    if avec_fig:
        blocs.append(F(fg.boite_groupes(groupes, y, g), f"Distribution de {y} selon {g} (losange = moyenne)"))
    blocs.append(T(norm, "Test de normalité de Shapiro-Wilk dans chaque groupe"))
    niv = int(round((1 - alpha) * 100))

    if k == 2:
        a, b = noms
        r = su.comparer_deux(groupes[a], groupes[b], alpha=alpha)
        reco, pv = r["reco"], r["p_reco"]
        ligne = {"t de Welch": "t de Welch (variances inégales)", "Mann-Whitney U": "Mann-Whitney U (non paramétrique)"}[reco]
        row = r["tests"].loc[ligne]
        stat = (f"t({_vf(row['ddl'])}) = {n2(row['Statistique'])}" if reco == "t de Welch"
                else f"U = {n2(row['Statistique'], 1)}")
        blocs += [P(f"Le test {reco} a été retenu car {r['motif']}."),
                  T(r["tests"], "Résultats des tests de comparaison de deux groupes"),
                  P(f"Résultat du test retenu : {stat}, {su.texte_p(pv)}.")]
        det = (f"La moyenne de {y} est de {n2(groupes[a].mean())} dans le groupe « {a} » et de "
               f"{n2(groupes[b].mean())} dans le groupe « {b} » (différence = {n2(r['diff'])} ; "
               f"IC à {niv} % : [{n2(r['ic'][0])} ; {n2(r['ic'][1])}]).")
        if reco != "t de Welch":
            det += (f" Les médianes sont respectivement de {n2(np.median(groupes[a]))} et "
                    f"{n2(np.median(groupes[b]))}.")
        effet = row["Taille d'effet"]
        if effet != "—":
            det += f" Taille d'effet : {_vf(effet)}"
            det += f", où X1 désigne le groupe « {a} »." if reco != "t de Welch" else "."
        blocs.append(P(det))
        test = reco
        sig = pv < alpha
        resume = (f"{y} diffère significativement entre les groupes de {g} ({su.texte_p(pv)})." if sig
                  else f"Aucune différence significative de {y} entre les groupes de {g} ({su.texte_p(pv)}).")
    else:
        r = su.comparer_k(groupes, alpha=alpha)
        reco, pv = r["reco"], r["p_reco"]
        ligne = {"ANOVA à un facteur": "ANOVA à un facteur", "ANOVA de Welch": "ANOVA de Welch",
                 "Kruskal-Wallis": "Kruskal-Wallis (non paramétrique)"}[reco]
        row = r["tests"].loc[ligne]
        stat = (f"H({row['ddl']}) = {n2(row['Statistique'])}" if reco == "Kruskal-Wallis"
                else f"F({_vf(row['ddl'])}) = {n2(row['Statistique'])}")
        blocs += [P(f"Le test {reco} a été retenu car {r['motif']}."),
                  T(r["tests"], f"Résultats des tests de comparaison de {k} groupes"),
                  P(f"Résultat du test retenu : {stat}, {su.texte_p(pv)}.")]
        effet = row["Taille d'effet"]
        if effet == "—":
            effet = r["tests"].loc["ANOVA à un facteur", "Taille d'effet"]
        blocs.append(P(f"Taille d'effet : {_vf(effet)}."))
        indic = (lambda v: np.median(v)) if reco == "Kruskal-Wallis" else (lambda v: np.mean(v))
        nom_indic = "médiane" if reco == "Kruskal-Wallis" else "moyenne"
        valeurs = {n: indic(groupes[n]) for n in noms}
        haut, bas = max(valeurs, key=valeurs.get), min(valeurs, key=valeurs.get)
        blocs.append(P(f"La {nom_indic} la plus élevée est observée dans le groupe « {haut} » ({n2(valeurs[haut])}) "
                       f"et la plus faible dans le groupe « {bas} » ({n2(valeurs[bas])})."))
        if pv < alpha:
            if r["posthoc"] == "Tukey HSD":
                ph = su.posthoc_tukey(groupes, alpha=alpha)
                nom_ph = "Tukey HSD"
            else:
                ph = su.posthoc_paires(groupes, test="welch" if reco == "ANOVA de Welch" else "mannwhitney",
                                       methode="holm", alpha=alpha)
                nom_ph = "comparaisons par paires avec correction de Holm"
            blocs.append(T(ph, f"Comparaisons post-hoc ({nom_ph})"))
            paires = [f"« {l['Groupe 1']} » et « {l['Groupe 2']} »" for _, l in ph.iterrows() if l["Significatif"] == "Oui"]
            if paires:
                blocs.append(P(f"Les comparaisons post-hoc ({nom_ph}) indiquent des différences significatives entre : "
                               + " ; ".join(paires) + "."))
            else:
                blocs.append(P("Aucune paire de groupes ne se distingue significativement après correction des comparaisons multiples."))
        test = reco
        sig = pv < alpha
        resume = (f"{y} diffère significativement selon {g} ({su.texte_p(pv)})." if sig
                  else f"Aucune différence significative de {y} selon {g} ({su.texte_p(pv)}).")
    blocs.append(P(_conclusion_p(pv, alpha)))
    return {"blocs": blocs, "p": pv, "test": test, "decision": _decision(pv, alpha), "resume": resume}


def _t_khi2(df, p, alpha, avec_fig):
    a, b = p["a"], p["b"]
    r = su.tester_khi2(df, a, b)
    ct, faible = r["effectifs"], r["pct_attendus_faibles"]
    min_att = float(r["attendus"].values.min())
    if r["table_2x2"] and min_att < 5:
        test, pv = "Test exact de Fisher", r["p_fisher"]
        motif = "le tableau est de dimension 2 × 2 et au moins un effectif théorique est inférieur à 5"
        stat = f"p exact = {su.formater_p(pv)}"
    else:
        test, pv = "Khi² de Pearson", r["p"]
        ligne = r["tests"].loc["Khi² de Pearson"]
        motif = f"{100 - faible:.0f} % des cellules ont un effectif théorique supérieur ou égal à 5".replace(".", ",")
        stat = f"χ²({ligne['ddl']}) = {n2(ligne['Statistique'])}, N = {r['n']}"
    dim = min(ct.shape) - 1
    seuils = {1: (0.1, 0.3, 0.5), 2: (0.07, 0.21, 0.35)}.get(dim, (0.06, 0.17, 0.29))
    force = su.force_effet(r["cramer_v"], seuils, ("négligeable", "faible", "modérée", "forte"))
    pct = r["pct_ligne"]
    ct_aff = ct.copy()
    pct_aff = pct.round(1)
    blocs = [P(f"L'association entre les variables qualitatives « {a} » (en lignes) et « {b} » (en colonnes) "
               f"est étudiée sur {r['n']} observations complètes."),
             T(ct_aff, f"Effectifs observés : {a} (lignes) × {b} (colonnes)"),
             T(pct_aff, f"Pourcentages en ligne (%) : répartition de {b} dans chaque modalité de {a}")]
    if avec_fig:
        blocs.append(F(fg.barres_empilees(pct, a, b), f"Répartition de {b} selon {a}"))
    blocs += [P(f"Le test retenu est le {test} car {motif}."),
              T(r["tests"], "Résultats des tests d'indépendance"),
              P(f"Résultat : {stat}, {su.texte_p(pv)}. Le V de Cramér vaut {n2(r['cramer_v'], 3)}, "
                f"ce qui correspond à une association {force}.")]
    if faible > 20:
        blocs.append(P(f"Attention : {n2(faible, 0)} % des cellules ont un effectif théorique inférieur à 5 "
                       f"(seuil usuel toléré : 20 %). Le résultat doit être interprété avec prudence ; "
                       f"regrouper des modalités peu fréquentes (onglet Recodage) peut améliorer la validité du test."))
    if pv < alpha:
        res = r["residus"]
        cellules = sorted(((abs(res.loc[i, j]), i, j, res.loc[i, j]) for i in res.index for j in res.columns), reverse=True)
        notables = [c for c in cellules if c[0] >= 2][:4]
        if notables:
            blocs.append(P("L'examen des résidus standardisés (|résidu| ≥ 2) montre que les combinaisons suivantes "
                           "s'écartent le plus de l'indépendance : " + " ; ".join(
                               f"{a} = {i} et {b} = {j} ({'sur-représentée' if v > 0 else 'sous-représentée'}, résidu = {n2(v)})"
                               for _, i, j, v in notables) + "."))
    blocs.append(P(_conclusion_p(pv, alpha)))
    sig = pv < alpha
    resume = (f"{a} et {b} sont significativement associées ({su.texte_p(pv)}, V de Cramér = {n2(r['cramer_v'], 2)})." if sig
              else f"Pas d'association significative entre {a} et {b} ({su.texte_p(pv)}).")
    return {"blocs": blocs, "p": pv, "test": test, "decision": _decision(pv, alpha), "resume": resume}


def _t_correlation(df, p, alpha, avec_fig):
    x, y, meth = p["x"], p["y"], p["methode"]
    _, _, _, lg = su.correlations(df, [x, y], methode=meth)
    if lg.empty or pd.isna(lg.iloc[0]["r"]):
        raise ValueError("corrélation non calculable (effectif insuffisant ou variable constante)")
    lig = lg.iloc[0]
    r, pv, n = float(lig["r"]), float(lig["p-value"]), int(lig["N"])
    nom = {"pearson": "de Pearson", "spearman": "de Spearman (rangs)", "kendall": "τ de Kendall"}[meth]
    symb = {"pearson": "r", "spearman": "ρ", "kendall": "τ"}[meth]
    force = "quasi nulle" if lig["Force"] == "nulle" else lig["Force"]
    tab = pd.DataFrame([{"Méthode": nom.replace("de ", "", 1).capitalize(), "Coefficient": n2(r, 3),
                         "p-value": pv, "N": n, "Intensité": force}])
    blocs = [P(f"La relation entre « {x} » et « {y} » est mesurée par le coefficient de corrélation {nom} "
               f"sur {n} observations complètes."), T(tab, f"Corrélation entre {x} et {y}")]
    if avec_fig:
        blocs.append(F(fg.nuage_reg(df[x], df[y], x, y), f"Nuage de points : {y} selon {x}"))
    sens = "positive : quand l'une augmente, l'autre tend à augmenter" if r > 0 else \
           "négative : quand l'une augmente, l'autre tend à diminuer"
    txt = f"Le coefficient vaut {symb} = {n2(r, 3)} ({su.texte_p(pv)}) : la relation est {sens}, d'intensité {force}."
    if meth == "pearson":
        txt += f" Environ {n2(100 * r * r, 1)} % de la variance est partagée entre les deux variables."
        d = df[[x, y]].apply(pd.to_numeric, errors="coerce").dropna()
        if n < 30 and not (su.est_normal(d[x], alpha) and su.est_normal(d[y], alpha)):
            txt += (" Les effectifs sont faibles et au moins une variable s'écarte de la normalité : "
                    "une corrélation de Spearman serait plus prudente.")
    blocs += [P(txt), P("Une corrélation mesure une association statistique et n'établit pas, à elle seule, une relation de cause à effet."),
              P(_conclusion_p(pv, alpha))]
    sig = pv < alpha
    resume = (f"{x} et {y} sont significativement corrélées ({symb} = {n2(r, 2)}, {su.texte_p(pv)})." if sig
              else f"Pas de corrélation significative entre {x} et {y} ({su.texte_p(pv)}).")
    return {"blocs": blocs, "p": pv, "test": f"Corrélation {nom}", "decision": _decision(pv, alpha), "resume": resume}


def _t_apparie(df, p, alpha, avec_fig):
    a, b = p["a"], p["b"]
    r = su.comparer_apparies(df[a], df[b], alpha=alpha)
    reco, pv, n = r["reco"], r["p_reco"], r["n"]
    ligne = {"t apparié": "t apparié", "Wilcoxon": "Wilcoxon (rangs signés)"}[reco]
    row = r["tests"].loc[ligne]
    stat = f"t({row['ddl']}) = {n2(row['Statistique'])}" if reco == "t apparié" else f"W = {n2(row['Statistique'], 1)}"
    niv = int(round((1 - alpha) * 100))
    blocs = [P(f"Les mesures « {a} » et « {b} » sont comparées sur {n} paires complètes.")]
    if avec_fig:
        blocs.append(F(fg.boite_appariees(df[a], df[b], a, b), f"Comparaison de {a} et {b}"))
    blocs += [P(f"Le test {reco} a été retenu car {r['motif']}."),
              T(r["tests"], "Résultats des tests sur données appariées"),
              P(f"Résultat : {stat}, {su.texte_p(pv)}. La différence moyenne ({a} − {b}) est de {n2(r['diff_moy'])} "
                f"(IC à {niv} % : [{n2(r['ic'][0])} ; {n2(r['ic'][1])}])."),
              P(_conclusion_p(pv, alpha))]
    sig = pv < alpha
    resume = (f"{a} et {b} diffèrent significativement ({su.texte_p(pv)})." if sig
              else f"Pas de différence significative entre {a} et {b} ({su.texte_p(pv)}).")
    return {"blocs": blocs, "p": pv, "test": reco, "decision": _decision(pv, alpha), "resume": resume}


def _t_normalite(df, p, alpha, avec_fig):
    v = p["v"]
    s = pd.to_numeric(df[v], errors="coerce").dropna()
    t = su.tests_normalite(s, alpha)
    if t.empty:
        raise ValueError("effectif insuffisant ou variable constante")
    ref = t[t["Test"] == "Shapiro-Wilk"]
    lig = ref.iloc[0] if not ref.empty else t.iloc[0]
    pv = float(lig["p-value"])
    d = su.decrire_numeriques(df, [v]).iloc[0]
    blocs = [P(f"La normalité de « {v} » est examinée sur {len(s)} valeurs valides."),
             T(t.set_index("Test"), f"Tests de normalité de {v}")]
    if avec_fig:
        blocs.append(F(fg.normalite(s, v), f"Distribution de {v} : histogramme et diagramme Q-Q"))
    txt = (f"Le test de référence est le test de {lig['Test']} (statistique = {n2(lig['Statistique'], 3)}, {su.texte_p(pv)}). "
           f"L'asymétrie vaut {n2(d['Asymétrie'])} et l'aplatissement {n2(d['Aplatissement'])}.")
    blocs += [P(txt), P(_conclusion_p(pv, alpha))]
    if pv < alpha and len(s) >= 30:
        blocs.append(P("Malgré cet écart, avec au moins 30 observations les tests sur les moyennes restent généralement "
                       "robustes (théorème central limite) ; un test non paramétrique reste une alternative prudente."))
    resume = (f"{v} s'écarte significativement de la normalité ({su.texte_p(pv)})." if pv < alpha
              else f"{v} est compatible avec la normalité ({su.texte_p(pv)}).")
    return {"blocs": blocs, "p": pv, "test": lig["Test"], "decision": _decision(pv, alpha), "resume": resume}


def _t_regression(df, p, alpha, avec_fig):
    y, xs = p["y"], p["xs"]
    cat = [c for c in xs if not pd.api.types.is_numeric_dtype(df[c])]
    modele, coef, res, diag = su.regression_ols(df, y, xs, categoriques=cat)
    pv = float(res["p_f"])
    k = int(modele.df_model)
    blocs = [P(f"La variable dépendante est « {y} ». Les variables explicatives sont : "
               + ", ".join(f"« {c} »" for c in xs) + (f" (qualitatives : {', '.join(cat)}, codées en indicatrices)" if cat else "")
               + f". Le modèle est estimé par les moindres carrés ordinaires sur {res['n']} observations complètes."),
             T(coef, f"Coefficients du modèle de régression de {y}"),
             P(f"Le modèle expliquerait {n2(100 * res['r2'], 1)} % de la variance de {y} (R² = {n2(res['r2'], 3)} ; "
               f"R² ajusté = {n2(res['r2_aj'], 3)}). Test global : F({_vf(res['ddl'])}) = {n2(res['f'])}, {su.texte_p(pv)}.")]
    sig = [(i, l) for i, l in coef.iterrows() if i != "Intercept" and l["p-value"] < alpha]
    if sig:
        blocs.append(P("Variables dont l'effet est significatif, toutes choses égales par ailleurs :"))
        blocs.append(L(f"{i} : coefficient = {n2(l['Coefficient'], 3)} ({su.texte_p(l['p-value'])}) — "
                       f"{'hausse' if l['Coefficient'] > 0 else 'baisse'} de {y} de {n2(abs(l['Coefficient']), 3)} "
                       f"par unité (ou par rapport à la modalité de référence pour une variable qualitative)"
                       for i, l in sig))
    else:
        blocs.append(P("Aucun coefficient pris isolément n'est significatif au seuil retenu."))
    if avec_fig:
        blocs.append(F(fg.diag_regression(modele), "Diagnostics du modèle : résidus et diagramme Q-Q des résidus"))
    diag_txt = []
    if "shapiro_p" in diag:
        ok = diag["shapiro_p"] >= alpha
        diag_txt.append(f"les résidus {'sont compatibles avec la normalité' if ok else 's’écartent de la normalité'} "
                        f"(Shapiro-Wilk, {su.texte_p(diag['shapiro_p'])})")
    if "bp_p" in diag and not pd.isna(diag["bp_p"]):
        ok = diag["bp_p"] >= alpha
        diag_txt.append(f"la variance des résidus {'est homogène' if ok else 'n’est pas homogène (hétéroscédasticité)'} "
                        f"(Breusch-Pagan, {su.texte_p(diag['bp_p'])})")
    if diag_txt:
        blocs.append(P("Vérification des hypothèses du modèle : " + " ; ".join(diag_txt) + "."))
    if res["n"] < 10 * max(k, 1):
        blocs.append(P(f"Attention : avec {res['n']} observations pour {k} paramètre(s) estimé(s), l'effectif est inférieur à "
                       f"la règle usuelle de 10 observations par paramètre ; les estimations peuvent être instables."))
    blocs.append(P(_conclusion_p(pv, alpha)))
    resume = (f"Le modèle explique significativement {y} (R² = {n2(res['r2'], 2)}, {su.texte_p(pv)})." if pv < alpha
              else f"Le modèle n'explique pas significativement {y} ({su.texte_p(pv)}).")
    return {"blocs": blocs, "p": pv, "test": "Régression linéaire (F global)", "decision": _decision(pv, alpha), "resume": resume}


def _t_fiabilite(df, p, alpha, avec_fig):
    items = p["items"]
    a, n, detail = su.cronbach(df, items)
    ok = a >= 0.7
    blocs = [P(f"La cohérence interne de l'échelle composée de {len(items)} items est évaluée par l'alpha de Cronbach "
               f"sur {n} répondants ayant répondu à tous les items."),
             T(detail, "Analyse des items de l'échelle")]
    if avec_fig:
        blocs.append(F(fg.barres_items(detail), "Corrélation item-total corrigée de chaque item"))
    blocs.append(P(f"L'alpha de Cronbach vaut {n2(a, 3)}, soit une fiabilité {su.lire_alpha(a)}. "
                   f"Le critère usuel (α ≥ 0,70) est {'atteint' if ok else 'non atteint'}."))
    faibles = [i for i, l in detail.iterrows() if l["Corrélation item-total corrigée"] < 0.3]
    ameliorent = [i for i, l in detail.iterrows() if l["Alpha si item supprimé"] > a + 0.01]
    if faibles:
        blocs.append(P("Items faiblement corrélés au reste de l'échelle (corrélation item-total < 0,30) : " + ", ".join(faibles) + "."))
    if ameliorent:
        blocs.append(P("La suppression de l'un des items suivants augmenterait l'alpha : " + ", ".join(ameliorent)
                       + ". Vérifiez aussi que tous les items sont orientés dans le même sens (inversion dans l'onglet Recodage)."))
    blocs.append(P("L'alpha de Cronbach est un indicateur de fiabilité et non un test d'hypothèse : la décision repose sur le seuil de 0,70."))
    resume = (f"L'échelle est fiable (α = {n2(a, 2)})." if ok else f"L'échelle n'atteint pas le seuil de fiabilité (α = {n2(a, 2)}).")
    return {"blocs": blocs, "p": None, "test": "Alpha de Cronbach",
            "decision": "Critère atteint" if ok else "Critère non atteint", "resume": resume}


EXECUTION = {"groupes": _t_groupes, "khi2": _t_khi2, "correlation": _t_correlation, "apparie": _t_apparie,
             "normalite": _t_normalite, "regression": _t_regression, "fiabilite": _t_fiabilite}


# ---------------------------------------------------------------- sections descriptives
def _bloc_descriptives(df, cols_num, cols_qual, avec_fig):
    blocs = []
    if cols_num:
        d = su.decrire_numeriques(df, cols_num)
        colonnes = ["N valides", "Manquants", "Moyenne", "Écart-type", "Min", "Q1", "Médiane", "Q3", "Max", "Asymétrie"]
        blocs += [H(2, "Variables quantitatives"),
                  P(f"Le tableau suivant présente les principaux indicateurs de {len(cols_num)} variable(s) quantitative(s)."),
                  T(d[colonnes], "Statistiques descriptives des variables quantitatives")]
        asym = [c for c in cols_num if pd.notna(d.loc[c, "Asymétrie"]) and abs(d.loc[c, "Asymétrie"]) > 1]
        if asym:
            blocs.append(P("Les variables " + ", ".join(f"« {c} »" for c in asym) + " présentent une forte asymétrie "
                           "(|asymétrie| > 1) : la moyenne y est sensible aux valeurs extrêmes, et les tests non paramétriques "
                           "ou une transformation (logarithme, racine) peuvent être préférables."))
        manq = [c for c in cols_num if d.loc[c, "Manquants"] > 0]
        if manq:
            blocs.append(P("Valeurs manquantes : " + ", ".join(f"« {c} » ({int(d.loc[c, 'Manquants'])}, "
                           f"{n2(d.loc[c, '% manquants'], 1)} %)" for c in manq) + "."))
        if avec_fig:
            for i in range(0, len(cols_num), 6):
                lot = cols_num[i:i + 6]
                blocs.append(F(fg.grille_histogrammes(df, lot), "Distribution des variables quantitatives : " + ", ".join(lot)))
    if cols_qual:
        blocs += [H(2, "Variables qualitatives")]
        phrases = []
        for c in cols_qual:
            f = su.frequences(df[c])
            valides = f.drop(index="Manquant", errors="ignore")
            if valides.empty:
                continue
            top = valides["Effectif"].idxmax()
            phrases.append(f"« {c} » : modalité la plus fréquente = « {top} » ({n2(valides.loc[top, '% valides'], 1)} % des réponses valides, "
                           f"{len(valides)} modalité(s))")
            blocs.append(T(valides[["Effectif", "% valides"]].head(12), f"Répartition de la variable {c}"))
        if phrases:
            blocs.append(L(phrases))
        if avec_fig:
            for i in range(0, len(cols_qual), 4):
                lot = cols_qual[i:i + 4]
                blocs.append(F(fg.grille_barres(df, lot), "Répartition des variables qualitatives : " + ", ".join(lot)))
    return blocs


# ---------------------------------------------------------------- construction du rapport
def construire_rapport(df, meta, hypotheses, options, alpha, recodages):
    """Exécute les tests et rédige le rapport ; retourne la liste de blocs."""
    avec_fig = options.get("graphiques", True)
    a_txt = n2(alpha, 2)
    blocs, num = [], 1

    # 1. Contexte
    blocs.append(H(1, f"{num}. Contexte et objectifs"))
    ctx = _paragraphes(meta.get("contexte")) + _paragraphes(meta.get("objectifs"))
    if ctx:
        blocs += ctx
    else:
        blocs.append(P(f"Ce rapport présente l'analyse statistique des données du fichier « {meta.get('fichier', '')} »."))
    if hypotheses:
        blocs.append(H(2, "Hypothèses à tester"))
        tab = pd.DataFrame([{"N°": i, "Analyse": titre_hypothese(h), "H0 (hypothèse nulle)": h["h0"], "H1 (hypothèse alternative)": h["h1"]}
                            for i, h in enumerate(hypotheses, 1)])
        blocs.append(T(tab, "Hypothèses statistiques testées"))
        for i, h in enumerate(hypotheses, 1):
            if h.get("enonce"):
                blocs.append(P(h["enonce"], gras=f"Hypothèse de recherche {i} : "))
    num += 1

    # 2. Données et méthodologie
    n, nv = df.shape
    manq = int(df.isna().sum().sum())
    blocs.append(H(1, f"{num}. Données et méthodologie"))
    blocs.append(P(f"Les données proviennent du fichier « {meta.get('fichier', '')} »"
                   + (f" (feuille « {meta['feuille']} »)" if meta.get("feuille") else "")
                   + f". Après nettoyage, elles comptent {n} observations et {nv} variables, avec "
                   f"{manq} valeur(s) manquante(s) ({n2(100 * manq / max(n * nv, 1), 1)} % des cellules)."))
    if options.get("inclure_recodages", True) and recodages:
        blocs.append(P("Les recodages suivants ont été appliqués avant l'analyse :"))
        blocs.append(L(f"{s['nom']} (à partir de {s['source']}) : {su.decrire_recodage(s)}" for s in recodages))
    blocs.append(P("Méthode statistique : tous les tests sont bilatéraux, au seuil de signification "
                   f"α = {a_txt}. Les valeurs manquantes sont exclues test par test. Le choix entre test paramétrique et "
                   "non paramétrique repose sur la normalité des données (test de Shapiro-Wilk pour les petits effectifs ; "
                   "au-delà de 30 observations par groupe, le théorème central limite est invoqué) et, pour plusieurs groupes, "
                   "sur l'homogénéité des variances (test de Levene). Les tailles d'effet sont rapportées en complément des p-values. "
                   "Les analyses ont été réalisées avec StatExcel (Python : pandas, SciPy, statsmodels)."))
    blocs.append(P(f"Règle de décision : si p < {a_txt}, l'hypothèse nulle H0 est rejetée ; sinon elle n'est pas rejetée."))
    num += 1

    # 3. Descriptives
    if options.get("descriptives", True) and (options.get("cols_num") or options.get("cols_qual")):
        blocs.append(H(1, f"{num}. Statistiques descriptives"))
        blocs += _bloc_descriptives(df, options.get("cols_num", []), options.get("cols_qual", []), avec_fig)
        num += 1

    # 4. Tests
    resultats = []
    if hypotheses:
        blocs.append(H(1, f"{num}. Résultats des tests d'hypothèses"))
        for i, h in enumerate(hypotheses, 1):
            blocs.append(H(2, f"{num}.{i} Hypothèse {i} : {titre_hypothese(h)}"))
            blocs.append(P(h["h0"], gras="H0 : "))
            blocs.append(P(h["h1"], gras="H1 : "))
            try:
                res = EXECUTION[h["type"]](df, h["params"], alpha, avec_fig)
            except Exception as e:
                motif = ("une des variables n'existe plus (recodage supprimé ou fichier changé)"
                         if isinstance(e, KeyError) else str(e))
                res = {"blocs": [P(f"Ce test n'a pas pu être réalisé : {motif}.")], "p": None, "test": "—",
                       "decision": "Non réalisé", "resume": f"Hypothèse {i} : test non réalisé ({motif})."}
            blocs += res["blocs"]
            resultats.append(res)
        num += 1

        # 5. Synthèse
        blocs.append(H(1, f"{num}. Synthèse et conclusion"))
        synth = pd.DataFrame([{"N°": i, "Analyse": titre_hypothese(h), "Test": r["test"], "p-value": r["p"], "Décision": r["decision"]}
                              for i, (h, r) in enumerate(zip(hypotheses, resultats), 1)])
        blocs.append(T(synth, "Synthèse des hypothèses testées"))
        avec_p = [(i, r) for i, r in enumerate(resultats, 1) if r["p"] is not None and not pd.isna(r["p"])]
        rejet = [i for i, r in avec_p if r["p"] < alpha]
        blocs.append(P(f"Sur {len(avec_p)} test(s) d'hypothèse, {len(rejet)} conduisent au rejet de H0 au seuil α = {a_txt}"
                       + (f" (hypothèses {', '.join(map(str, rejet))})." if rejet else ".")))
        blocs.append(L(r["resume"] for r in resultats))
        if len(avec_p) >= 2:
            m = len(avec_p)
            seuil = alpha / m
            tient = [i for i, r in avec_p if r["p"] < seuil]
            blocs.append(P(f"Plusieurs tests ({m}) ayant été réalisés sur les mêmes données, le risque cumulé de faux positif augmente. "
                           f"À titre de contrôle, avec la correction de Bonferroni le seuil par test serait de {n2(seuil, 4)} : "
                           + (f"les hypothèses {', '.join(map(str, tient))} restent significatives." if tient
                              else "aucune hypothèse ne reste significative.")))
        num += 1

    # Limites
    blocs.append(H(1, f"{num}. Limites et précautions"))
    blocs.append(L([
        "Un résultat statistiquement significatif ne mesure pas l'importance pratique d'un effet : il faut lire les tailles d'effet et les intervalles de confiance.",
        "Les associations et corrélations observées ne prouvent pas de relation de cause à effet.",
        "Un résultat non significatif ne démontre pas l'absence d'effet ; il peut résulter d'un effectif insuffisant.",
        "La validité des tests dépend de la qualité des données et de la représentativité de l'échantillon.",
    ]))
    return blocs


# ---------------------------------------------------------------- mise en forme des tableaux
def _col_p(nom):
    n = str(nom).lower()
    return n == "p" or "p-value" in n or n.startswith("p ") or n.endswith("(p)")


def _cellule(col, x):
    if isinstance(x, (list, tuple, np.ndarray)):
        return str(x)
    if x is None or pd.isna(x):
        return "—"
    if isinstance(x, (bool, np.bool_)):
        return "Oui" if x else "Non"
    if isinstance(x, (int, np.integer)):
        return str(int(x))
    if isinstance(x, (float, np.floating)):
        x = float(x)
        if _col_p(col):
            return su.formater_p(x)
        if x.is_integer() and abs(x) < 1e6:
            return str(int(x))
        if abs(x) >= 1000:
            return f"{x:,.0f}".replace(",", " ")
        if abs(x) >= 100:
            return n2(x, 1)
        return n2(x, 4 if 0 < abs(x) < 0.01 else 2)
    return str(x)


def formater_table(df):
    """DataFrame -> DataFrame de texte, nombres à la française, index inclus comme première colonne."""
    d = df.copy()
    if not isinstance(d.index, pd.RangeIndex):
        d = d.reset_index()
        if str(d.columns[0]) == "index":
            d = d.rename(columns={"index": ""})
    d.columns = [str(c) for c in d.columns]
    return pd.DataFrame({c: [_cellule(c, v) for v in d[c]] for c in d.columns})


# ---------------------------------------------------------------- export Word
def exporter_docx(blocs, meta):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    NAVY = RGBColor(0x1F, 0x3A, 0x5F)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin = sec.bottom_margin = Cm(2.3)

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(att), "Calibri")
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "fr-FR")
    rpr.append(lang)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for nom, taille in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        st_ = doc.styles[nom]
        st_.font.name, st_.font.size, st_.font.bold, st_.font.color.rgb = "Calibri", Pt(taille), True, NAVY
        st_.paragraph_format.space_before = Pt(16 if nom == "Heading 1" else 10)
        st_.paragraph_format.space_after = Pt(6)
        r_ = st_.element.get_or_add_rPr()
        f_ = r_.find(qn("w:rFonts"))
        if f_ is None:
            f_ = OxmlElement("w:rFonts")
            r_.append(f_)
        for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            f_.set(qn(att), "Calibri")
        for att in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if f_.get(qn(att)) is not None:
                del f_.attrib[qn(att)]

    def champ(par, instr):
        run = par.add_run()
        for typ, txt in (("begin", None), (None, instr), ("end", None)):
            if typ:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), typ)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = txt
            run._r.append(el)

    # pied de page : auteur + numéro de page
    pied = sec.footer.paragraphs[0]
    pied.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = pied.add_run(f"{meta.get('auteur', '')} — StatExcel — page ")
    r0.font.size = Pt(9)
    champ(pied, "PAGE")
    for r_ in pied.runs:
        r_.font.size = Pt(9)
        r_.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    # page de titre
    for _ in range(6):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ = t.add_run(meta.get("titre") or "Rapport d'analyse statistique")
    r_.bold, r_.font.size, r_.font.color.rgb = True, Pt(26), NAVY
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ = s.add_run("Rapport d'analyse statistique")
    r_.font.size, r_.font.color.rgb = Pt(14), RGBColor(0x6B, 0x72, 0x80)
    for _ in range(5):
        doc.add_paragraph()
    for ligne, gras in ((meta.get("auteur", ""), True), (meta.get("organisation", ""), False), (date_fr(), False)):
        if ligne:
            p_ = doc.add_paragraph()
            p_.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_ = p_.add_run(ligne)
            r_.bold, r_.font.size = gras, Pt(13)
    doc.add_page_break()

    def ombrer(cell, couleur):
        pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), couleur)
        pr.append(shd)

    n_tab = n_fig = 0
    for b in blocs:
        typ = b["type"]
        if typ in ("h1", "h2", "h3"):
            doc.add_heading(b["texte"], level=int(typ[1]))
        elif typ == "p":
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if b.get("gras"):
                par.add_run(b["gras"]).bold = True
            par.add_run(b["texte"])
        elif typ == "liste":
            for it in b["items"]:
                par = doc.add_paragraph(it, style="List Bullet")
                par.paragraph_format.space_after = Pt(3)
        elif typ == "table":
            n_tab += 1
            cap = doc.add_paragraph()
            cap.paragraph_format.keep_with_next = True
            cap.paragraph_format.space_before = Pt(8)
            r_ = cap.add_run(f"Tableau {n_tab} : ")
            r_.bold = True
            r_.font.size = Pt(9.5)
            r_ = cap.add_run(b["legende"])
            r_.italic, r_.font.size = True, Pt(9.5)
            d = formater_table(b["df"])
            taille = 9 if d.shape[1] <= 6 else 8 if d.shape[1] <= 9 else 7.5
            tab = doc.add_table(rows=1, cols=d.shape[1])
            tab.style = "Table Grid"
            tab.alignment = WD_TABLE_ALIGNMENT.CENTER
            tab.autofit = False
            longueurs = [min(max([len(str(c))] + [len(x) for x in d[c]]), 38) + 3 for c in d.columns]
            largeurs = [Cm(16.4 * l / sum(longueurs)) for l in longueurs]
            for j, c in enumerate(d.columns):
                cell = tab.rows[0].cells[j]
                ombrer(cell, "1F3A5F")
                par = cell.paragraphs[0]
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.space_after = Pt(0)
                r_ = par.add_run(c)
                r_.bold, r_.font.size, r_.font.color.rgb = True, Pt(taille), RGBColor(0xFF, 0xFF, 0xFF)
            for i in range(len(d)):
                cells = tab.add_row().cells
                for j, c in enumerate(d.columns):
                    if i % 2 == 1:
                        ombrer(cells[j], "EEF2F7")
                    par = cells[j].paragraphs[0]
                    par.paragraph_format.space_after = Pt(0)
                    texte = d.iloc[i, j]
                    long = len(texte) > 28
                    par.alignment = WD_ALIGN_PARAGRAPH.LEFT if (j == 0 or long) else WD_ALIGN_PARAGRAPH.CENTER
                    r_ = par.add_run(texte)
                    r_.font.size = Pt(taille)
                    if j == 0:
                        r_.bold = True
            for ri, row in enumerate(tab.rows):
                trpr = row._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:cantSplit"))
                if ri == 0:
                    trpr.append(OxmlElement("w:tblHeader"))
                for ci, cell in enumerate(row.cells):
                    cell.width = largeurs[ci]
                    if len(tab.rows) <= 16 and ri < len(tab.rows) - 1:
                        for par in cell.paragraphs:
                            par.paragraph_format.keep_with_next = True
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif typ == "fig":
            n_fig += 1
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.keep_with_next = True
            from PIL import Image
            pw, ph = Image.open(io.BytesIO(b["png"])).size
            largeur = min(15.5, 9.0 * pw / ph)
            par.add_run().add_picture(io.BytesIO(b["png"]), width=Cm(largeur))
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_ = cap.add_run(f"Figure {n_fig} : ")
            r_.bold, r_.font.size = True, Pt(9.5)
            r_ = cap.add_run(b["legende"])
            r_.italic, r_.font.size = True, Pt(9.5)

    tampon = io.BytesIO()
    doc.save(tampon)
    return tampon.getvalue()
