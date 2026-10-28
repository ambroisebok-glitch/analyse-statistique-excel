"""
Étudiants Learn (version Gemini) – cours (PDF/Word) -> résumé approfondi, QCM par chapitre (jusqu'à 100),
assistant IA, comptes utilisateurs et exports.
Auteur : BOKOUMBO Tibetobe Ambroise
Contact : ambroisebok@gmail.com | +228 91942948
"""
import hashlib
import hmac
import html
import io
import json
import os
import random
import re
import secrets
import sqlite3
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from functools import lru_cache

import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from google import genai
from google.genai import types
from pypdf import PdfReader

# ----------------------------------------------------------------------------
# Configuration
# ----------------------------------------------------------------------------
MODEL = "gemini-3.8-flash"
MAX_CHARS = 700_000                 # le cours est lu en entier (≈ 175 000 tokens)
MAX_PDF_OCTETS = 18 * 1024 * 1024   # taille max d'un PDF envoyé directement à Gemini
LOT_QCM = 20                        # questions générées par appel
MAX_QCM = 100
LETTRES = "ABCD"
DB_PATH = os.getenv("DB_PATH", "revisepro.db")
PAGES = ["📄 Mon cours", "📝 Résumé", "✅ QCM", "💬 Assistant", "🗂️ Historique", "🤝 Services & contact"]

AUTEUR = "BOKOUMBO Tibetobe Ambroise"
EMAIL = "ambroisebok@gmail.com"
TEL_REP = ("+228 96286244", "+22896286244")   # répétitions + démarches
TEL_AMB = ("+228 91942948", "+22891942948")   # contact principal
APP_URL = "https://etudiantslearn.streamlit.app/"
SEO_TITRE = "Étudiants Learn – Résumés de cours, QCM et assistant IA"
SEO_DESC = ("Étudiants Learn aide les étudiants à réviser leurs cours PDF ou Word grâce à des résumés "
            "approfondis, des QCM corrigés jusqu'à 100 questions et un assistant IA pour comprendre les cours "
            "et préparer les examens.")

st.set_page_config(
    page_title=SEO_TITRE, page_icon="🎓", layout="wide", initial_sidebar_state="expanded",
    menu_items={"About": f"**Étudiants Learn** — conçu par {AUTEUR} · {EMAIL} · {TEL_AMB[0]}"})

st.markdown(
    """
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
.stApp { font-family: 'Inter', sans-serif; background: linear-gradient(180deg,#F8F9FF 0%,#FFFFFF 320px); }
#MainMenu, footer { visibility: hidden; }
header[data-testid="stHeader"] { background: transparent; }
iframe[height="0"] { display:none; }
.block-container { padding-top: 1.6rem; padding-bottom: 7rem; max-width: 1100px; animation: fondu .4s ease; }
@keyframes fondu { from {opacity:0; transform:translateY(6px);} to {opacity:1; transform:none;} }

/* Barre latérale */
[data-testid="stSidebar"] { background: #FFFFFF; border-right: 1px solid #E5E7EB; }
.marque { display:flex; align-items:center; gap:.65rem; margin:.3rem 0 .6rem 0; }
.marque .logo { width:40px; height:40px; border-radius:12px; display:flex; align-items:center;
  justify-content:center; font-size:1.3rem; background:linear-gradient(135deg,#4F46E5,#7C3AED);
  box-shadow:0 4px 12px rgba(79,70,229,.35); }
.marque b { font-size:1.25rem; font-weight:800; color:#1E1B4B; letter-spacing:-.02em; }
[data-testid="stSidebar"] [role="radiogroup"] label { padding:.5rem .75rem; border-radius:10px;
  transition: background .15s; width:100%; }
[data-testid="stSidebar"] [role="radiogroup"] label:hover { background:#F1F2FF; }
[data-testid="stSidebar"] [role="radiogroup"] label:has(input:checked) { background:#EEF2FF;
  font-weight:600; color:#4338CA; }
.promo { background:linear-gradient(135deg,#EEF2FF,#F5F3FF); border:1px solid #E0E7FF; border-radius:14px;
  padding:.8rem .9rem; font-size:.82rem; color:#3730A3; line-height:1.45; }
.promo b { color:#1E1B4B; }

/* Bandeau et cartes */
.hero { position:relative; overflow:hidden; background: linear-gradient(135deg,#4338CA 0%,#7C3AED 60%,#A855F7 100%);
  color:#fff; padding: 1.7rem 2rem; border-radius: 20px; margin-bottom: 1.2rem;
  box-shadow: 0 12px 30px rgba(79,70,229,.28); }
.hero::after { content:""; position:absolute; right:-60px; top:-70px; width:240px; height:240px; border-radius:50%;
  background: radial-gradient(circle, rgba(255,255,255,.28), rgba(255,255,255,0) 70%); }
.hero h1 { color:#fff; font-size: 2rem; margin:0; padding:0; font-weight:800; letter-spacing:-.02em; }
.hero p { color:#E0E7FF; margin:.4rem 0 0 0; font-size:1.02rem; max-width: 720px; }
.cartes { display:grid; grid-template-columns:repeat(auto-fit,minmax(220px,1fr)); gap:.9rem; margin:.4rem 0 1.2rem 0; }
.carte { border:1px solid #E5E7EB; border-radius:16px; padding:1.1rem 1.2rem; background:#fff;
  box-shadow: 0 1px 3px rgba(0,0,0,.05); transition: transform .15s, box-shadow .15s; }
.carte:hover { transform: translateY(-3px); box-shadow: 0 10px 24px rgba(79,70,229,.12); }
.carte .ico { font-size:1.6rem; width:46px; height:46px; display:flex; align-items:center; justify-content:center;
  background:#EEF2FF; border-radius:12px; margin-bottom:.6rem; }
.carte h4 { margin:.1rem 0 .3rem 0; font-size:1.02rem; color:#1E1B4B; }
.carte p { margin:0; color:#4B5563; font-size:.9rem; line-height:1.5; }
.etapes { display:flex; flex-wrap:wrap; gap:.5rem; margin-bottom:1rem; }
.etape { padding:.35rem .8rem; border-radius:999px; font-size:.82rem; font-weight:600;
  background:#F3F4F6; color:#6B7280; border:1px solid #E5E7EB; }
.etape.ok { background:#ECFDF5; color:#047857; border-color:#A7F3D0; }
.btn { display:inline-block; padding:.55rem 1rem; border-radius:10px; font-weight:600; font-size:.9rem;
  text-decoration:none !important; background:#4F46E5; color:#fff !important; margin:.25rem .3rem .25rem 0; }
.btn:hover { background:#4338CA; }
.btn.wa { background:#16A34A; } .btn.wa:hover { background:#15803D; }
.btn.clair { background:#EEF2FF; color:#4338CA !important; }

/* Composants Streamlit */
.stButton>button, .stDownloadButton>button { border-radius:12px; font-weight:600; transition: all .15s; }
.stButton>button[kind="primary"] { background:linear-gradient(135deg,#4F46E5,#7C3AED); border:none;
  box-shadow:0 4px 12px rgba(79,70,229,.3); }
.stButton>button[kind="primary"]:hover { transform:translateY(-1px); box-shadow:0 8px 18px rgba(79,70,229,.38); }
div[data-testid="stMetric"] { background:#fff; border:1px solid #E5E7EB; border-radius:14px;
  padding:.8rem 1rem; box-shadow:0 1px 3px rgba(0,0,0,.04); }
div[data-testid="stExpander"] { border-radius:14px; border:1px solid #E5E7EB; background:#fff; }
div[data-testid="stForm"] { border-radius:16px; border:1px solid #E5E7EB; background:#fff; }

/* Barre fixe en bas */
.barre-pied { position:fixed; left:0; right:0; bottom:0; z-index:99; background:#1E1B4B; color:#C7D2FE;
  padding:.5rem 1rem; text-align:center; font-size:.8rem; line-height:1.55;
  box-shadow:0 -4px 16px rgba(0,0,0,.15); }
.barre-pied a { color:#fff !important; text-decoration:none; font-weight:600; }
.barre-pied a:hover { text-decoration:underline; }
.barre-pied .sep { opacity:.4; margin:0 .5rem; }
.barre-pied .acc { color:#FDE68A; }
@media (max-width: 640px) { .barre-pied { font-size:.72rem; } .hero h1 { font-size:1.5rem; } }
</style>
""",
    unsafe_allow_html=True,
)


def injecter_seo():
    """Ajoute les balises SEO dans le <head> de la page (meilleur effort : Streamlit est une
    application dynamique, le site vitrine statique reste indispensable pour un bon référencement)."""
    if st.session_state.get("seo_ok"):
        return
    st.session_state.seo_ok = True
    donnees = {
        "titre": SEO_TITRE, "desc": SEO_DESC, "url": APP_URL,
        "json": {
            "@context": "https://schema.org", "@type": "WebApplication", "name": "Étudiants Learn",
            "url": APP_URL, "description": SEO_DESC, "inLanguage": "fr",
            "applicationCategory": "EducationalApplication", "operatingSystem": "Web",
            "author": {"@type": "Person", "name": AUTEUR, "email": EMAIL, "telephone": TEL_AMB[0]},
        },
    }
    js = """<script>
    try {
      const D = __DATA__, d = window.parent.document;
      d.documentElement.lang = "fr";
      const meta = (a, k, v) => { let m = d.head.querySelector('meta[' + a + '="' + k + '"]');
        if (!m) { m = d.createElement("meta"); m.setAttribute(a, k); d.head.appendChild(m); }
        m.setAttribute("content", v); };
      meta("name", "description", D.desc);
      meta("name", "keywords", "réviser cours, résumé de cours IA, QCM en ligne, quiz, assistant IA étudiant, " +
        "répétiteur Togo, démarches en ligne Togo, passeport, nationalité, casier judiciaire");
      meta("name", "author", "BOKOUMBO Tibetobe Ambroise");
      meta("name", "robots", "index, follow");
      meta("property", "og:type", "website"); meta("property", "og:title", D.titre);
      meta("property", "og:description", D.desc); meta("property", "og:url", D.url);
      meta("property", "og:locale", "fr_FR");
      meta("name", "twitter:card", "summary"); meta("name", "twitter:title", D.titre);
      let c = d.head.querySelector('link[rel="canonical"]');
      if (!c) { c = d.createElement("link"); c.rel = "canonical"; d.head.appendChild(c); }
      c.href = D.url;
      if (!d.getElementById("ld-revisepro")) { const s = d.createElement("script");
        s.type = "application/ld+json"; s.id = "ld-revisepro"; s.textContent = JSON.stringify(D.json);
        d.head.appendChild(s); }
    } catch (e) {}
    </script>""".replace("__DATA__", json.dumps(donnees))
    if hasattr(st, "iframe"):          # Streamlit récent
        st.iframe(js, height=1)
    else:                              # anciennes versions
        import streamlit.components.v1 as components
        components.html(js, height=0)


injecter_seo()


# ----------------------------------------------------------------------------
# Base de données (SQLite) et comptes utilisateurs
# ----------------------------------------------------------------------------
def maintenant() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def sql(requete, params=(), fetch=None):
    c = sqlite3.connect(DB_PATH)
    c.row_factory = sqlite3.Row
    try:
        cur = c.execute(requete, params)
        c.commit()
        if fetch == "one":
            r = cur.fetchone()
            return dict(r) if r else None
        if fetch == "all":
            return [dict(r) for r in cur.fetchall()]
        return cur.lastrowid
    finally:
        c.close()


@st.cache_resource
def init_db():
    c = sqlite3.connect(DB_PATH)
    c.executescript(
        """
        CREATE TABLE IF NOT EXISTS users(
            id INTEGER PRIMARY KEY AUTOINCREMENT, email TEXT UNIQUE, nom TEXT,
            sel TEXT, hash TEXT, cree TEXT);
        CREATE TABLE IF NOT EXISTS cours(
            id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, nom TEXT,
            texte TEXT, pdf BLOB, chapitres TEXT, cree TEXT);
        CREATE TABLE IF NOT EXISTS resumes(
            id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, cours_id INTEGER,
            titre TEXT, contenu TEXT, cree TEXT);
        CREATE TABLE IF NOT EXISTS qcms(
            id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, cours_id INTEGER,
            titre TEXT, data TEXT, cree TEXT);
        CREATE TABLE IF NOT EXISTS tentatives(
            id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, qcm_id INTEGER,
            score INTEGER, total INTEGER, cree TEXT);
        CREATE TABLE IF NOT EXISTS chats(
            cours_id INTEGER, user_id INTEGER, messages TEXT,
            PRIMARY KEY(cours_id, user_id));
        """
    )
    c.commit()
    c.close()
    return True


def hacher(mdp: str, sel: str) -> str:
    return hashlib.pbkdf2_hmac("sha256", mdp.encode(), bytes.fromhex(sel), 200_000).hex()


def creer_compte(nom: str, email: str, mdp: str):
    email = email.strip().lower()
    if len(nom.strip()) < 2:
        return None, "Entre ton nom."
    if not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email):
        return None, "Adresse e-mail invalide."
    if len(mdp) < 8:
        return None, "Le mot de passe doit contenir au moins 8 caractères."
    sel = secrets.token_hex(16)
    try:
        i = sql("INSERT INTO users(email,nom,sel,hash,cree) VALUES(?,?,?,?,?)",
                (email, nom.strip(), sel, hacher(mdp, sel), maintenant()))
    except sqlite3.IntegrityError:
        return None, "Un compte existe déjà avec cet e-mail."
    return {"id": i, "nom": nom.strip(), "email": email}, None


def connecter(email: str, mdp: str):
    u = sql("SELECT * FROM users WHERE email=?", (email.strip().lower(),), fetch="one")
    if not u or not hmac.compare_digest(u["hash"], hacher(mdp, u["sel"])):
        return None, "E-mail ou mot de passe incorrect."
    return {"id": u["id"], "nom": u["nom"], "email": u["email"]}, None


def uid():
    u = st.session_state.get("user")
    return u["id"] if u else None


def enregistrer_cours(c: dict) -> int:
    return sql("INSERT INTO cours(user_id,nom,texte,pdf,chapitres,cree) VALUES(?,?,?,?,?,?)",
               (uid(), c["nom"], c["texte"], c.get("pdf"), json.dumps(c["chapitres"]), maintenant()))


def charger_cours(cid: int):
    r = sql("SELECT * FROM cours WHERE id=? AND user_id=?", (cid, uid()), fetch="one")
    if not r:
        return None
    return {"id": r["id"], "nom": r["nom"], "texte": r["texte"], "pdf": r["pdf"],
            "chapitres": json.loads(r["chapitres"]), "vision": bool(r["pdf"]),
            "pages": 0}


def supprimer_cours(cid: int):
    sql("DELETE FROM tentatives WHERE qcm_id IN (SELECT id FROM qcms WHERE cours_id=? AND user_id=?)",
        (cid, uid()))
    for t in ("resumes", "qcms", "chats"):
        sql(f"DELETE FROM {t} WHERE cours_id=? AND user_id=?", (cid, uid()))
    sql("DELETE FROM cours WHERE id=? AND user_id=?", (cid, uid()))


def charger_chat(cid):
    if not (uid() and cid):
        return []
    r = sql("SELECT messages FROM chats WHERE cours_id=? AND user_id=?", (cid, uid()), fetch="one")
    return json.loads(r["messages"]) if r else []


def sauver_chat(cid, msgs):
    if uid() and cid:
        sql("INSERT OR REPLACE INTO chats(cours_id,user_id,messages) VALUES(?,?,?)",
            (cid, uid(), json.dumps(msgs)))


# ----------------------------------------------------------------------------
# Lecture des fichiers
# ----------------------------------------------------------------------------
def lire_fichier(fichier) -> dict:
    nom = fichier.name
    data = fichier.getvalue()
    pdf, pages = None, 0
    if nom.lower().endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        pages = len(reader.pages)
        texte = "\n".join((p.extract_text() or "") for p in reader.pages)
        if len(data) <= MAX_PDF_OCTETS:
            pdf = data
    elif nom.lower().endswith(".docx"):
        doc = Document(io.BytesIO(data))
        lignes = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                lignes.append(" | ".join(c.text for c in row.cells))
        texte = "\n".join(lignes)
    else:
        raise ValueError("Format non pris en charge (PDF ou .docx uniquement).")
    texte = re.sub(r"[ \t]+", " ", texte)
    texte = re.sub(r"\n{3,}", "\n\n", texte).strip()
    return {"nom": nom, "texte": texte, "pdf": pdf, "pages": pages, "vision": pdf is not None}


# ----------------------------------------------------------------------------
# Gemini
# ----------------------------------------------------------------------------
@st.cache_data(ttl=3600, show_spinner=False)
def modeles_disponibles(cle: str) -> list:
    """Modèle préféré d'abord, puis les autres modèles 'flash' du compte."""
    noms = [MODEL]
    try:
        client = genai.Client(api_key=cle)
        for m in client.models.list():
            n = m.name.replace("models/", "")
            actions = getattr(m, "supported_actions", None) or ["generateContent"]
            exclus = ("image", "tts", "live", "audio", "embedding", "robotics")
            if ("flash" in n and "generateContent" in actions
                    and not any(x in n for x in exclus) and n not in noms):
                noms.append(n)
    except Exception:
        pass
    return noms[:5]


def _temporaire(msg: str) -> bool:
    return any(m in msg for m in ("503", "UNAVAILABLE", "429", "RESOURCE_EXHAUSTED",
                                  "500", "INTERNAL", "overloaded", "DEADLINE"))


def parts_cours(cours: dict) -> list:
    if cours.get("pdf") and cours.get("vision"):
        return [types.Part.from_bytes(data=cours["pdf"], mime_type="application/pdf")]
    return [types.Part.from_text(text="COURS :\n" + cours["texte"][:MAX_CHARS])]


def requete(cours: dict, prompt: str) -> list:
    return [types.Content(role="user",
                          parts=parts_cours(cours) + [types.Part.from_text(text=prompt)])]


REGLAGES = {"sans_reflexion": True, "max_tokens": None}  # mémorise ce que le modèle accepte


def generer(contenu, cle: str, system=None, max_tokens=16000, temperature=0.4,
            reflexion=True) -> str:
    client = genai.Client(api_key=cle)

    def config():
        kw = dict(system_instruction=system, temperature=temperature,
                  max_output_tokens=min(max_tokens, REGLAGES["max_tokens"] or max_tokens))
        if not reflexion and REGLAGES["sans_reflexion"]:
            kw["thinking_config"] = types.ThinkingConfig(thinking_budget=0)
        return types.GenerateContentConfig(**kw)

    derniere = None
    for modele in modeles_disponibles(cle):
        for essai in range(3):
            try:
                rep = client.models.generate_content(model=modele, contents=contenu, config=config())
                if rep.text and rep.text.strip():
                    return rep.text
                derniere = RuntimeError("Réponse vide")
            except Exception as e:
                derniere = e
                msg = str(e)
                if "400" in msg or "INVALID_ARGUMENT" in msg:
                    # Réglage refusé par le modèle : on retire d'abord la "réflexion",
                    # puis on plafonne la longueur de sortie, sinon c'est une vraie erreur.
                    if not reflexion and REGLAGES["sans_reflexion"]:
                        REGLAGES["sans_reflexion"] = False
                        continue
                    if (REGLAGES["max_tokens"] or max_tokens) > 8192:
                        REGLAGES["max_tokens"] = 8192
                        continue
                    raise
                if "404" in msg or "NOT_FOUND" in msg:
                    break  # modèle indisponible : passer au suivant
                if _temporaire(msg):
                    time.sleep(3 * (essai + 1))
                    continue
                raise
    raise RuntimeError("Gemini est surchargé ou son quota est atteint. "
                       f"Réessaie dans quelques minutes. ({derniere})")


def flux_chat(contenu, cle: str, system: str):
    """Génère la réponse en streaming (avec repli sur d'autres modèles)."""
    client = genai.Client(api_key=cle)
    cfg = types.GenerateContentConfig(system_instruction=system, temperature=0.5,
                                      max_output_tokens=16000)
    derniere = None
    for modele in modeles_disponibles(cle):
        deja = False
        try:
            for morceau in client.models.generate_content_stream(
                    model=modele, contents=contenu, config=cfg):
                if morceau.text:
                    deja = True
                    yield morceau.text
            return
        except Exception as e:
            derniere = e
            if deja:
                raise
            msg = str(e)
            if "404" in msg or "NOT_FOUND" in msg or _temporaire(msg):
                time.sleep(2)
                continue
            raise
    raise RuntimeError(f"Gemini est indisponible pour le moment. ({derniere})")


def normaliser_math(t: str) -> str:
    """Convertit \\( \\) et \\[ \\] en $ et $$ (rendu LaTeX de Streamlit)."""
    t = re.sub(r"\\\[(.+?)\\\]", lambda m: "\n$$" + m.group(1).strip() + "$$\n", t, flags=re.S)
    t = re.sub(r"\\\((.+?)\\\)", lambda m: "$" + m.group(1).strip() + "$", t, flags=re.S)
    return t


REGLE_MATHS = ("Pour toutes les formules mathématiques, utilise LaTeX : $...$ en ligne et $$...$$ "
               "pour les formules sur leur propre ligne. N'utilise jamais \\( \\) ni \\[ \\].")


# ----------------------------------------------------------------------------
# Chapitres
# ----------------------------------------------------------------------------
def detecter_chapitres(cours: dict, cle: str) -> list:
    prompt = f"""Analyse la structure du cours fourni et liste ses chapitres (ou grandes parties) DANS L'ORDRE,
en couvrant TOUT le document du début à la fin. Si le cours n'a pas de chapitres explicites, découpe-le
en parties thématiques cohérentes (entre 3 et 12). Maximum 40 entrées.
Réponds UNIQUEMENT avec une ligne par chapitre, au format exact :
CHAPITRE | titre du chapitre | description en une phrase des notions couvertes
Aucun autre texte."""
    brut = generer(requete(cours, prompt), cle, max_tokens=6000, temperature=0.1, reflexion=False)
    chap = []
    for ligne in brut.splitlines():
        m = re.match(r"\s*CHAPITRE\s*\|\s*(.+?)\s*\|\s*(.*)", ligne)
        if m:
            chap.append({"titre": m.group(1).strip(" *#"), "desc": m.group(2).strip()})
    return chap[:40] or [{"titre": "Cours complet", "desc": ""}]


def portee_texte(c: dict) -> str:
    if c["titre"] == "Cours complet":
        return "l'ensemble du cours, en couvrant toutes ses parties de façon équilibrée"
    return f"le chapitre « {c['titre']} » ({c['desc']})" if c["desc"] else f"le chapitre « {c['titre']} »"


# ----------------------------------------------------------------------------
# Résumé approfondi
# ----------------------------------------------------------------------------
NIVEAUX = {
    "Essentiel": ("synthétique : environ une page par chapitre, uniquement l'essentiel "
                  "(définitions, formules, résultats à retenir)", 8000),
    "Détaillé": ("détaillée : environ deux à trois pages par chapitre, avec explications, "
                 "formules et exemples", 16000),
    "Approfondi": ("très approfondie et EXHAUSTIVE : reprends TOUT le contenu du chapitre sans rien omettre, "
                   "avec démonstrations, exemples résolus étape par étape (calculs intermédiaires détaillés), "
                   "remarques, cas particuliers et applications. Plus c'est complet et précis, mieux c'est",
                   32000),
}


def generer_resume(cours, chapitres, niveau, cle, progres=None) -> dict:
    consigne, max_tok = NIVEAUX[niveau]
    liste = "\n".join(f"- {c['titre']}" for c in chapitres)
    modeles_disponibles(cle)  # préchauffe le cache (avant les threads)

    def apercu_fn():
        prompt = f"""À partir du cours fourni, rédige la « Vue d'ensemble » d'une fiche de révision portant sur :
{liste}
Format strict : la première ligne est « TITRE: <titre du cours> ». Ensuite, en Markdown
(titres de niveau ### uniquement) :
### Résumé de synthèse (3 à 4 paragraphes denses)
### Plan et objectifs d'apprentissage
### Notions clés à maîtriser (liste à puces)
### Glossaire (15 à 25 termes, au format : **terme** : définition)
Rédige en français. {REGLE_MATHS}"""
        return generer(requete(cours, prompt), cle, max_tokens=12000, reflexion=False)

    def fiche_fn(c):
        prompt = f"""Rédige la fiche de révision du chapitre suivant du cours fourni : {portee_texte(c)}.
Cette fiche doit être {consigne}.
Structure en Markdown, avec des titres de niveau ### (jamais # ni ##) :
### Idée générale
### Notions et définitions   (termes importants en **gras**)
### Théorèmes, propriétés et formules   (énoncés complets avec leurs conditions d'application)
### Démonstrations et justifications importantes   (étapes détaillées)
### Méthodes et techniques
### Exemples résolus   (ceux du cours, résolus pas à pas, puis un ou deux exemples types supplémentaires)
### Pièges et erreurs fréquentes
### À retenir   (mini fiche mémo)
Omets une rubrique seulement si elle n'a aucun sens pour ce chapitre. Utilise des tableaux Markdown quand
c'est utile. Rédige en français, sans phrase d'introduction du type « Voici ». {REGLE_MATHS}"""
        return generer(requete(cours, prompt), cle, max_tokens=max_tok)

    total = len(chapitres) + 1
    brut_apercu, mds, echecs, fait = "", [None] * len(chapitres), 0, 0
    with ThreadPoolExecutor(max_workers=min(4, total)) as ex:
        futs = {ex.submit(apercu_fn): -1}
        for k, c in enumerate(chapitres):
            futs[ex.submit(fiche_fn, c)] = k
        for f in as_completed(futs):
            k = futs[f]
            try:
                res = f.result()
            except Exception as e:
                res = None
                if k >= 0:
                    echecs += 1
                    mds[k] = f"*⚠️ Cette section n'a pas pu être générée ({e}). Relance le résumé.*"
                else:
                    brut_apercu = ""
            if res is not None:
                if k < 0:
                    brut_apercu = res
                else:
                    mds[k] = normaliser_math(res.strip())
            fait += 1
            if progres:
                progres(fait / total, f"{fait}/{total} parties terminées")

    titre = cours["nom"].rsplit(".", 1)[0]
    m = re.match(r"\s*TITRE\s*:\s*(.+)", brut_apercu)
    if m:
        titre = m.group(1).strip(" *#")
        brut_apercu = brut_apercu[m.end():]
    sections = [{"titre": c["titre"], "md": mds[k]} for k, c in enumerate(chapitres)]
    return {"titre": titre, "niveau": niveau, "apercu": normaliser_math(brut_apercu.strip()),
            "sections": sections, "echecs": echecs, "cree": maintenant()}


def resume_en_markdown(r: dict) -> str:
    out = [f"# {r['titre']}", "", r["apercu"], ""]
    for s in r["sections"]:
        out += [f"## {s['titre']}", "", s["md"], ""]
    return "\n".join(out)


# ----------------------------------------------------------------------------
# Exports : HTML (impression PDF avec formules), Word (formules en images)
# ----------------------------------------------------------------------------
def slug(t: str) -> str:
    return re.sub(r"[^a-zA-Z0-9]+", "_", t).strip("_")[:50] or "resume"


def md_vers_html(md_texte: str, titre: str) -> str:
    import markdown
    maths = []

    def garde(m):
        maths.append(m.group(0))
        return f"@@MATH{len(maths) - 1}@@"

    t = re.sub(r"\$\$.+?\$\$|\$[^$\n]+?\$", garde, md_texte, flags=re.S)
    corps = markdown.markdown(t, extensions=["tables", "fenced_code", "sane_lists"])
    corps = re.sub(r"@@MATH(\d+)@@", lambda m: html.escape(maths[int(m.group(1))], quote=False), corps)
    modele = """<!DOCTYPE html><html lang="fr"><head><meta charset="utf-8">
<title>__TITRE__</title>
<script>window.MathJax={tex:{inlineMath:[['$','$']],displayMath:[['$$','$$']]}};</script>
<script async src="https://cdn.jsdelivr.net/npm/mathjax@3/es5/tex-mml-chtml.js"></script>
<style>
body{font-family:Georgia,'Times New Roman',serif;max-width:820px;margin:2rem auto;padding:0 1.2rem;
line-height:1.65;color:#1F2937}
h1{color:#4F46E5;border-bottom:3px solid #4F46E5;padding-bottom:.3rem}
h2{color:#4F46E5;margin-top:2.2rem;border-bottom:1px solid #E5E7EB;padding-bottom:.2rem;
page-break-before:auto}
h3{color:#374151;margin-top:1.4rem}
table{border-collapse:collapse;margin:1rem 0}th,td{border:1px solid #D1D5DB;padding:.4rem .7rem}
th{background:#EEF2FF}code{background:#F3F4F6;padding:.1rem .3rem;border-radius:4px}
.note{color:#6B7280;font-size:.85rem;border-top:1px solid #E5E7EB;margin-top:3rem;padding-top:.6rem}
@media print{body{margin:0;max-width:none}.noprint{display:none}h2{page-break-before:always}}
</style></head><body>
<p class="noprint" style="background:#EEF2FF;padding:.6rem 1rem;border-radius:8px;font-family:sans-serif">
💡 Pour obtenir un PDF : <b>Ctrl+P</b> (ou Partager → Imprimer) puis « Enregistrer au format PDF ».</p>
__CORPS__
<p class="note">Généré par Étudiants Learn · __DATE__</p></body></html>"""
    return (modele.replace("__TITRE__", html.escape(titre)).replace("__CORPS__", corps)
            .replace("__DATE__", maintenant()))


DPI_FORMULE = 300


def _nettoyer_latex(e: str) -> str:
    e = e.replace("\\dfrac", "\\frac").replace("\\tfrac", "\\frac").replace("\\displaystyle", "")
    e = re.sub(r"\\(?:text|operatorname|mathrm)\{", r"\\mathrm{", e)
    e = re.sub(r"\\(?:,|;|!|:|quad|qquad)", " ", e)
    e = e.replace("\\left.", "").replace("\\right.", "").replace("\\Bigg", "").replace("\\Big", "")
    return e.strip()


@lru_cache(maxsize=2048)
def formule_png(expr: str, taille: int):
    """Rend une formule LaTeX en PNG avec matplotlib. Retourne (octets, largeur_px) ou None."""
    try:
        import matplotlib
        matplotlib.use("Agg")
        from matplotlib.figure import Figure
        from PIL import Image
        matplotlib.rcParams["mathtext.fontset"] = "cm"
        fig = Figure(figsize=(0.1, 0.1), dpi=DPI_FORMULE)
        fig.text(0, 0, f"${_nettoyer_latex(expr)}$", fontsize=taille)
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=DPI_FORMULE, transparent=True,
                    bbox_inches="tight", pad_inches=0.03)
        data = buf.getvalue()
        return data, Image.open(io.BytesIO(data)).size[0]
    except Exception:
        return None


RE_INLINE = re.compile(r"(\$\$.+?\$\$|\$[^$\n]+?\$|\*\*.+?\*\*|(?<!\*)\*[^*\n]+?\*(?!\*)|`[^`\n]+`)", re.S)


def _formule(par, expr, taille):
    rendu = formule_png(expr, taille)
    if rendu:
        par.add_run().add_picture(io.BytesIO(rendu[0]), width=Inches(rendu[1] / DPI_FORMULE))
    else:
        r = par.add_run(expr)
        r.italic = True


def _inline(par, texte):
    for m in RE_INLINE.split(texte):
        if not m:
            continue
        if m.startswith("$$"):
            _formule(par, m.strip("$").strip(), 12)
        elif m.startswith("$"):
            _formule(par, m.strip("$").strip(), 11)
        elif m.startswith("**"):
            par.add_run(m[2:-2]).bold = True
        elif m.startswith("`"):
            r = par.add_run(m[1:-1])
            r.font.name = "Consolas"
        elif m.startswith("*") and len(m) > 2:
            par.add_run(m[1:-1]).italic = True
        else:
            par.add_run(m)


def _md_dans_docx(doc, md_texte: str):
    t = re.sub(r"\$\$(.+?)\$\$", lambda m: "$$" + " ".join(m.group(1).split()) + "$$", md_texte, flags=re.S)
    lignes, i = t.splitlines(), 0
    while i < len(lignes):
        l = lignes[i].rstrip()
        s = l.strip()
        if not s:
            i += 1
            continue
        if s.startswith("|"):
            bloc = []
            while i < len(lignes) and lignes[i].strip().startswith("|"):
                bloc.append(lignes[i].strip())
                i += 1
            rows = [[c.strip() for c in b.strip("|").split("|")] for b in bloc
                    if not re.match(r"^\|?[\s:\-|]+\|?$", b)]
            if rows:
                n = max(len(r) for r in rows)
                tab = doc.add_table(rows=len(rows), cols=n)
                try:
                    tab.style = "Light Grid Accent 1"
                except Exception:
                    tab.style = "Table Grid"
                for a, r in enumerate(rows):
                    for b, c in enumerate(r):
                        cell = tab.cell(a, b)
                        cell.paragraphs[0].text = ""
                        _inline(cell.paragraphs[0], c)
                doc.add_paragraph()
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", s)
        if m:
            niveau = max(1, min(len(m.group(1)) - 1, 4))
            doc.add_heading(m.group(2).strip(" *"), level=niveau)
        elif s.startswith("$$") and s.endswith("$$") and len(s) > 4:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _formule(p, s[2:-2].strip(), 13)
        elif re.match(r"^[-*•]\s+", s):
            indent = len(l) - len(l.lstrip())
            p = doc.add_paragraph(style="List Bullet 2" if indent >= 2 else "List Bullet")
            _inline(p, re.sub(r"^[-*•]\s+", "", s))
        elif re.match(r"^\d+[.)]\s+", s):
            p = doc.add_paragraph(style="List Number")
            _inline(p, re.sub(r"^\d+[.)]\s+", "", s))
        else:
            _inline(doc.add_paragraph(), s)
        i += 1


def resume_en_docx(r: dict) -> bytes:
    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.9)
        sec.top_margin = sec.bottom_margin = Inches(0.8)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    for nom in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        try:
            doc.styles[nom].font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)
        except Exception:
            pass
    t = doc.add_heading(r["titre"], level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(f"Fiche de révision ({r.get('niveau', '')}) · {r.get('cree', '')}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _md_dans_docx(doc, r["apercu"])
    for s in r["sections"]:
        doc.add_page_break()
        doc.add_heading(s["titre"], level=1)
        _md_dans_docx(doc, s["md"])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------
# QCM
# ----------------------------------------------------------------------------
PATRON_QCM = re.compile(
    r"Question\s*:\s*(.+?)\n\s*A\s*[\).]\s*(.+?)\n\s*B\s*[\).]\s*(.+?)\n\s*C\s*[\).]\s*(.+?)\n"
    r"\s*D\s*[\).]\s*(.+?)\n\s*R[ée]ponse\s*:\s*\**\s*([ABCD])\b\**[ \t]*(?:\n\s*Explication\s*:\s*(.+))?",
    re.S | re.I)


def parser_qcm(brut: str, chap: str) -> list:
    out = []
    for bloc in brut.split("@@Q")[1:]:
        m = PATRON_QCM.search(bloc.strip() + "\n")
        if not m:
            continue
        opts = [normaliser_math(x.strip()) for x in m.groups()[1:5]]
        bonne = LETTRES.index(m.group(6).upper())
        ordre = list(range(4))
        random.shuffle(ordre)  # évite le biais de position de la bonne réponse
        out.append({"q": normaliser_math(m.group(1).strip()),
                    "o": [opts[k] for k in ordre], "r": ordre.index(bonne),
                    "e": normaliser_math((m.group(7) or "").strip()), "chap": chap})
    return out


DIFFICULTES = {
    "Facile": "facile (restitution des définitions et résultats de base)",
    "Moyen": "moyenne (compréhension et application directe)",
    "Difficile": "difficile (raisonnement, calculs, cas pièges)",
    "Mixte": "variée (environ 30 % faciles, 40 % moyennes, 30 % difficiles)",
}


def generer_qcm(cours, cibles, nb, difficulte, calculs, melanger, cle, progres=None) -> list:
    modeles_disponibles(cle)  # préchauffe le cache (avant les threads)
    # Répartition équilibrée des questions entre les chapitres ciblés
    base, reste = divmod(nb, len(cibles))
    bonus = set(random.sample(range(len(cibles)), reste))
    taches = [(c, base + (1 if i in bonus else 0)) for i, c in enumerate(cibles)]
    taches = [(c, n) for c, n in taches if n > 0]

    # Regroupe les petits chapitres : moins d'appels = plus rapide
    groupes, cour, somme = [], [], 0
    for c, n in taches:
        cour.append((c, n))
        somme += n
        if somme >= 6:
            groupes.append(cour)
            cour, somme = [], 0
    if cour:
        if groupes:
            groupes[-1] += cour
        else:
            groupes.append(cour)
    jobs = []
    for g in groupes:
        n = sum(x[1] for x in g)
        if len(g) == 1:
            jobs.append((g[0][0]["titre"], portee_texte(g[0][0]), n))
        else:
            noms = [x[0]["titre"] for x in g]
            jobs.append((" · ".join(noms)[:90],
                         "les chapitres suivants, en répartissant équitablement les questions : "
                         + " ; ".join(f"« {x} »" for x in noms), n))

    def produire(label, portee, n):
        recues, tentatives, derniere = [], 0, None
        while len(recues) < n and tentatives < 3:
            tentatives += 1
            lot = min(LOT_QCM, n - len(recues))
            evite = "; ".join(q["q"][:70].replace("\n", " ") for q in recues[-40:]) or "aucune pour l'instant"
            prompt = f"""Tu es un professeur expérimenté qui prépare un examen. À partir du cours fourni,
rédige {lot} questions à choix multiples portant UNIQUEMENT sur {portee}.
Règles :
- Questions fidèles au contenu du cours, précises et autonomes ; 4 options plausibles, UNE seule bonne réponse.
- Difficulté : {DIFFICULTES[difficulte]}.
- Varie les types : définition, compréhension, application, piège classique.
{"- Inclus des questions de calcul / d'application numérique avec le calcul justifié dans l'explication." if calculs else ""}
- L'explication justifie la bonne réponse (et le calcul si besoin) en 1 à 3 phrases.
- Ne reprends pas ces questions déjà posées : {evite}
- {REGLE_MATHS}
FORMAT STRICT (aucun JSON, aucun autre texte), répété pour chaque question :
@@Q
Question: ...
A) ...
B) ...
C) ...
D) ...
Reponse: une seule lettre parmi A, B, C, D
Explication: ..."""
            try:
                brut = generer(requete(cours, prompt), cle, max_tokens=12000,
                               temperature=0.7, reflexion=False)
            except Exception as e:
                derniere = e
                continue
            recues += parser_qcm(brut, label)[:lot]
        if not recues and derniere:
            raise derniere
        return recues

    par_job, erreurs, fait = {}, [], 0
    with ThreadPoolExecutor(max_workers=min(4, len(jobs))) as ex:
        futs = {ex.submit(produire, *j): k for k, j in enumerate(jobs)}
        for f in as_completed(futs):
            k = futs[f]
            try:
                par_job[k] = f.result()
            except Exception as e:
                erreurs.append(e)
                par_job[k] = []
            fait += jobs[k][2]
            if progres:
                progres(min(fait / nb, 1.0),
                        f"{sum(len(v) for v in par_job.values())} / {nb} questions générées")
    if erreurs and not any(par_job.values()):
        raise erreurs[0]

    deja, resultat = set(), []
    for k in range(len(jobs)):  # on garde l'ordre du cours
        for q in par_job.get(k, []):
            cle_q = re.sub(r"\W+", "", q["q"].lower())[:80]
            if cle_q not in deja:
                deja.add(cle_q)
                resultat.append(q)
    if melanger:
        random.shuffle(resultat)
    return resultat


def sauver_qcm(titre, questions):
    c = st.session_state.get("cours") or {}
    if uid():
        return sql("INSERT INTO qcms(user_id,cours_id,titre,data,cree) VALUES(?,?,?,?,?)",
                   (uid(), c.get("id"), titre, json.dumps(questions), maintenant()))
    return None


# ----------------------------------------------------------------------------
# Interface : composants communs
# ----------------------------------------------------------------------------
def entete(titre, sous=""):
    st.markdown(f'<div class="hero"><h1>{titre}</h1><p>{sous}</p></div>', unsafe_allow_html=True)


def cle_api():
    cle = os.getenv("GEMINI_API_KEY")
    if not cle:
        try:
            cle = st.secrets.get("GEMINI_API_KEY")
        except Exception:
            cle = None
    if not cle:
        cle = st.text_input("Clé API Gemini", type="password",
                            help="Ou définis la variable GEMINI_API_KEY (secrets Streamlit).")
    return cle


def pied():
    st.markdown(
        f'''<div class="barre-pied">
        📚 <b>Étudiants Learn</b> · Conçu par <b>{AUTEUR}</b><span class="sep">|</span>
        ✉️ <a href="mailto:{EMAIL}">{EMAIL}</a><span class="sep">|</span>
        📞 <a href="tel:{TEL_AMB[1]}">{TEL_AMB[0]}</a> · <a href="tel:{TEL_REP[1]}">{TEL_REP[0]}</a><br>
        <span class="acc">🎓 Répétitions : <a href="tel:{TEL_REP[1]}">{TEL_REP[0]}</a></span><span class="sep">|</span>
        <span class="acc">🗂️ Démarches en ligne (nationalité, passeport, casier judiciaire) &amp; traitement de données :
        <a href="tel:{TEL_REP[1]}">{TEL_REP[0]}</a> ou <a href="tel:{TEL_AMB[1]}">{TEL_AMB[0]}</a></span>
        </div>''', unsafe_allow_html=True)


def boutons_contact(numero, message="Bonjour, je vous contacte depuis Étudiants Learn."):
    from urllib.parse import quote
    return (f'<a class="btn" href="tel:{numero[1]}">📞 Appeler {numero[0]}</a>'
            f'<a class="btn wa" target="_blank" href="https://wa.me/{numero[1].lstrip("+")}?text={quote(message)}">'
            f'💬 WhatsApp</a>')


def bloc_services():
    st.markdown(
        f'''<div class="cartes">
        <div class="carte"><div class="ico">🎓</div><h4>Répétitions</h4>
        <p>Besoin d'un accompagnement personnalisé pour réviser ? Contactez le <b>{TEL_REP[0]}</b>.</p></div>
        <div class="carte"><div class="ico">🗂️</div><h4>Démarches en ligne</h4>
        <p>Nous vous accompagnons dans vos demandes en ligne : nationalité, passeport, casier judiciaire.</p></div>
        <div class="carte"><div class="ico">📊</div><h4>Traitement de données</h4>
        <p>Saisie, nettoyage, analyse et mise en forme de vos données.<br>
        📞 <b>{TEL_REP[0]}</b> ou <b>{TEL_AMB[0]}</b></p></div></div>''', unsafe_allow_html=True)


def page_services():
    entete("🤝 Services & contact", "Au-delà de l'application, nous vous accompagnons concrètement.")
    bloc_services()
    st.subheader("🎓 Répétitions")
    st.write("Pour des besoins de répétition (soutien scolaire, préparation d'examens, méthodes de travail), "
             "vous pouvez contacter le **" + TEL_REP[0] + "**.")
    st.markdown(boutons_contact(TEL_REP, "Bonjour, je souhaite des informations sur les répétitions."),
                unsafe_allow_html=True)
    st.subheader("🗂️ Demandes en ligne & traitement de données")
    st.write("Nous vous accompagnons dans vos demandes en ligne (**nationalité, passeport, casier judiciaire**) "
             "ainsi que dans le **traitement de vos données**.")
    st.markdown(boutons_contact(TEL_REP, "Bonjour, je souhaite être accompagné(e) pour une démarche en ligne.")
                + boutons_contact(TEL_AMB, "Bonjour, je souhaite être accompagné(e) pour une démarche en ligne."),
                unsafe_allow_html=True)
    st.caption("Service d'accompagnement privé, non affilié à l'administration. "
               "Les délais et décisions dépendent des services officiels.")
    st.subheader("✉️ Nous écrire")
    st.markdown(f'<a class="btn clair" href="mailto:{EMAIL}">✉️ {EMAIL}</a>', unsafe_allow_html=True)


def etapes():
    ok = [bool(st.session_state.get("cours")), bool(st.session_state.get("resume")),
          bool(st.session_state.get("qcm")), bool(st.session_state.get("chat"))]
    noms = ["1 · Charger le cours", "2 · Résumé", "3 · QCM", "4 · Assistant"]
    st.markdown('<div class="etapes">' + "".join(
        f'<span class="etape {"ok" if o else ""}">{"✓ " if o else ""}{n}</span>' for o, n in zip(ok, noms))
        + "</div>", unsafe_allow_html=True)


def besoin_cours():
    if not st.session_state.get("cours"):
        st.info("📄 Commence par charger un cours dans l'onglet **Mon cours**.")
        return None
    return st.session_state.cours


def reset_travail(cours):
    st.session_state.cours = cours
    st.session_state.resume = None
    st.session_state.qcm = None
    st.session_state.valide = False
    st.session_state.chat = charger_chat(cours.get("id"))


# ----------------------------------------------------------------------------
# Pages
# ----------------------------------------------------------------------------
def page_connexion():
    entete("📚 Étudiants Learn", "Transforme tes cours PDF ou Word en résumés approfondis, QCM par chapitre "
           "et assistant IA pour réviser plus vite et mieux.")
    st.markdown(
        '''<div class="cartes">
        <div class="carte"><div class="ico">📝</div><h4>Résumés approfondis</h4>
        <p>Une fiche par chapitre : définitions, formules, démonstrations, exemples. Export Word, PDF, Markdown.</p></div>
        <div class="carte"><div class="ico">✅</div><h4>QCM jusqu'à 100 questions</h4>
        <p>Sur tout le cours ou par chapitre, avec correction détaillée et note sur 20.</p></div>
        <div class="carte"><div class="ico">💬</div><h4>Assistant IA</h4>
        <p>Pose tes questions sur ton cours, demande des exercices corrigés et des explications.</p></div>
        <div class="carte"><div class="ico">🧮</div><h4>Maths et formules</h4>
        <p>Les formules LaTeX s'affichent proprement, dans l'application comme dans les exports.</p></div></div>''',
        unsafe_allow_html=True)
    _, centre, _ = st.columns([1, 1.6, 1])
    with centre:
        t1, t2, t3 = st.tabs(["Connexion", "Créer un compte", "Invité"])
        with t1:
            with st.form("login"):
                email = st.text_input("E-mail")
                mdp = st.text_input("Mot de passe", type="password")
                if st.form_submit_button("Se connecter", type="primary", use_container_width=True):
                    u, err = connecter(email, mdp)
                    if err:
                        st.error(err)
                    else:
                        st.session_state.user = u
                        st.rerun()
        with t2:
            with st.form("signup"):
                nom = st.text_input("Nom complet")
                email = st.text_input("E-mail", key="se")
                mdp = st.text_input("Mot de passe (8 caractères min.)", type="password", key="sm")
                if st.form_submit_button("Créer mon compte", type="primary", use_container_width=True):
                    u, err = creer_compte(nom, email, mdp)
                    if err:
                        st.error(err)
                    else:
                        st.session_state.user = u
                        st.rerun()
        with t3:
            st.write("Utilise l'application sans compte. Tes cours, résumés et QCM "
                     "ne seront **pas** sauvegardés.")
            if st.button("Continuer en invité", use_container_width=True):
                st.session_state.user = {"id": None, "nom": "Invité", "email": ""}
                st.rerun()
    st.write("")
    bloc_services()
    pied()


def page_cours(cle):
    entete("📄 Mon cours", "Dépose un PDF ou un Word : le cours est lu en entier et découpé en chapitres.")
    etapes()
    fichier = st.file_uploader("Fichier du cours", type=["pdf", "docx"])
    if st.button("🔍 Analyser le cours", type="primary", disabled=not fichier):
        if not cle:
            st.error("Ajoute ta clé API Gemini.")
        else:
            try:
                with st.status("Analyse du cours…", expanded=True) as s:
                    st.write("Lecture du fichier…")
                    cours = lire_fichier(fichier)
                    if len(cours["texte"]) < 80 and not cours["pdf"]:
                        raise ValueError("Aucun texte lisible (document vide ou PDF scanné trop volumineux).")
                    if len(cours["texte"]) > MAX_CHARS:
                        st.write("⚠️ Cours très long : seule la partie initiale sera exploitée en mode texte.")
                    st.write("Détection des chapitres…")
                    cours["chapitres"] = detecter_chapitres(cours, cle)
                    if uid():
                        st.write("Enregistrement dans ta bibliothèque…")
                        cours["id"] = enregistrer_cours(cours)
                    reset_travail(cours)
                    s.update(label=f"Cours prêt ✅ — {len(cours['chapitres'])} chapitre(s)", state="complete")
            except Exception as e:
                st.error(f"Erreur : {e}")

    cours = st.session_state.get("cours")
    if cours:
        st.subheader(cours["nom"])
        c1, c2, c3 = st.columns(3)
        c1.metric("Pages", cours.get("pages") or "—")
        c2.metric("Mots", f"{len(cours['texte'].split()):,}".replace(",", " "))
        c3.metric("Chapitres", len(cours["chapitres"]))
        if cours.get("pdf"):
            cours["vision"] = st.toggle(
                "Lecture visuelle du PDF (formules, schémas, scans) — recommandé",
                value=cours.get("vision", True),
                help="Gemini lit directement les pages du PDF. Désactive pour un mode texte plus rapide.")
        with st.expander("📚 Structure détectée", expanded=True):
            for i, c in enumerate(cours["chapitres"], 1):
                st.markdown(f"**{i}. {c['titre']}**" + (f" — {c['desc']}" if c["desc"] else ""))

    if uid():
        st.divider()
        st.subheader("🗃️ Ma bibliothèque")
        lst = sql("SELECT id, nom, cree FROM cours WHERE user_id=? ORDER BY id DESC", (uid(),), fetch="all")
        if not lst:
            st.caption("Aucun cours enregistré pour l'instant.")
        for c in lst:
            a, b, d = st.columns([5, 1.3, 1.3])
            a.markdown(f"**{c['nom']}**  \n<small>{c['cree']}</small>", unsafe_allow_html=True)
            if b.button("Ouvrir", key=f"o{c['id']}", use_container_width=True):
                reset_travail(charger_cours(c["id"]))
                st.rerun()
            if d.button("🗑️", key=f"s{c['id']}", use_container_width=True):
                supprimer_cours(c["id"])
                if (st.session_state.get("cours") or {}).get("id") == c["id"]:
                    st.session_state.cours = None
                st.rerun()
    else:
        st.info("Crée un compte pour sauvegarder tes cours, résumés et notes.")


def afficher_resume(r: dict, cle_export: str):
    st.header(r["titre"])
    st.caption(f"Niveau : {r.get('niveau', '')} · généré le {r.get('cree', '')}")
    md = resume_en_markdown(r)
    c1, c2, c3 = st.columns(3)
    c1.download_button("⬇️ HTML (→ PDF)", md_vers_html(md, r["titre"]).encode("utf-8"),
                       file_name=f"{slug(r['titre'])}.html", mime="text/html", use_container_width=True,
                       help="Ouvre le fichier puis Ctrl+P → Enregistrer en PDF (formules incluses).")
    c2.download_button("⬇️ Markdown", md.encode("utf-8"), file_name=f"{slug(r['titre'])}.md",
                       mime="text/markdown", use_container_width=True)
    if c3.button("📄 Préparer le Word (.docx)", use_container_width=True, key=f"pw{cle_export}"):
        with st.spinner("Création du document Word (rendu des formules)…"):
            st.session_state.docx = (cle_export, resume_en_docx(r))
    prete = st.session_state.get("docx")
    if prete and prete[0] == cle_export:
        st.download_button("✅ Télécharger le Word", prete[1], file_name=f"{slug(r['titre'])}.docx",
                           mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                           type="primary", use_container_width=True)
    st.divider()
    st.markdown(r["apercu"])
    st.divider()
    st.subheader("Fiches par chapitre")
    for i, s in enumerate(r["sections"]):
        with st.expander(f"{i + 1}. {s['titre']}", expanded=(i == 0 and len(r["sections"]) == 1)):
            st.markdown(s["md"])


def page_resume(cle):
    entete("📝 Résumé", "Une fiche de révision complète, chapitre par chapitre, avec formules et exemples.")
    cours = besoin_cours()
    if not cours:
        return
    titres = [c["titre"] for c in cours["chapitres"]]
    choix = st.multiselect("Chapitres à résumer", titres, default=titres)
    niveau = st.select_slider("Niveau de détail", list(NIVEAUX), value="Approfondi")
    if st.button("✨ Générer le résumé", type="primary", disabled=not choix):
        if not cle:
            st.error("Ajoute ta clé API Gemini.")
        else:
            cibles = [c for c in cours["chapitres"] if c["titre"] in choix]
            barre = st.progress(0.0, text="Démarrage…")
            try:
                r = generer_resume(cours, cibles, niveau, cle,
                                   lambda p, t: barre.progress(min(p, 1.0), text=t))
                if uid():
                    r["id"] = sql("INSERT INTO resumes(user_id,cours_id,titre,contenu,cree) VALUES(?,?,?,?,?)",
                                  (uid(), cours.get("id"), r["titre"], json.dumps(r), maintenant()))
                st.session_state.resume = r
                barre.empty()
                if r["echecs"]:
                    st.warning(f"{r['echecs']} section(s) non générée(s) : relance pour les compléter.")
            except Exception as e:
                barre.empty()
                st.error(f"Erreur : {e}")
    r = st.session_state.get("resume")
    if r:
        st.divider()
        afficher_resume(r, str(r.get("id") or r["cree"]))


def page_qcm(cle):
    entete("✅ QCM", "Jusqu'à 100 questions, sur tout le cours ou chapitre par chapitre, avec correction détaillée.")
    cours = besoin_cours()
    if not cours:
        return
    titres = [c["titre"] for c in cours["chapitres"]]

    with st.expander("⚙️ Paramètres du QCM", expanded=not st.session_state.get("qcm")):
        mode = st.radio("Portée", ["Cours entier (tous les chapitres, équilibré)", "Par chapitre(s)"],
                        horizontal=True)
        if mode.startswith("Par"):
            sel = st.multiselect("Chapitres", titres, default=titres[:1])
        else:
            sel = titres
        c1, c2 = st.columns(2)
        nb = c1.slider("Nombre de questions", 5, MAX_QCM, 20, step=5)
        diff = c2.select_slider("Difficulté", list(DIFFICULTES), value="Mixte")
        c3, c4 = st.columns(2)
        calculs = c3.toggle("Inclure des questions de calcul / d'application", value=True)
        melange = c4.toggle("Mélanger l'ordre des questions", value=False)
        go = st.button(f"🎯 Générer {nb} questions", type="primary", disabled=not sel)

    if go:
        if not cle:
            st.error("Ajoute ta clé API Gemini.")
        else:
            cibles = [c for c in cours["chapitres"] if c["titre"] in sel]
            barre = st.progress(0.0, text="Démarrage…")
            try:
                qs = generer_qcm(cours, cibles, nb, diff, calculs, melange, cle,
                                 lambda p, t: barre.progress(p, text=t))
                barre.empty()
                if not qs:
                    st.error("Aucune question valide générée. Réessaie.")
                else:
                    if len(qs) < nb:
                        st.warning(f"{len(qs)} questions valides générées sur {nb} demandées.")
                    portee = "cours entier" if len(sel) == len(titres) else ", ".join(sel)[:80]
                    titre = f"QCM {len(qs)} q. — {portee}"
                    st.session_state.qcm = {"id": sauver_qcm(titre, qs), "titre": titre, "questions": qs}
                    st.session_state.valide = False
                    st.session_state.version = st.session_state.get("version", 0) + 1
                    st.rerun()
            except Exception as e:
                barre.empty()
                st.error(f"Erreur : {e}")

    qcm = st.session_state.get("qcm")
    if not qcm:
        st.info("Règle les paramètres ci-dessus puis clique sur « Générer ».")
        return
    qs, v = qcm["questions"], st.session_state.get("version", 0)
    st.subheader(qcm["titre"])

    if not st.session_state.get("valide"):
        st.caption(f"{len(qs)} questions · une seule bonne réponse par question.")
        with st.form(f"form_{v}"):
            for i, q in enumerate(qs):
                st.markdown(f"**{i + 1}.** {q['q']}")
                if q.get("chap") and q["chap"] != "Cours complet":
                    st.caption(f"📚 {q['chap']}")
                st.radio("Réponse", range(4), format_func=lambda j, q=q: f"{LETTRES[j]}. {q['o'][j]}",
                         index=None, key=f"q{v}_{i}", label_visibility="collapsed")
                st.divider()
            ok = st.form_submit_button("✅ Valider et voir la correction", type="primary")
        if ok:
            reps = [st.session_state.get(f"q{v}_{i}") for i in range(len(qs))]
            st.session_state.reponses = reps
            st.session_state.valide = True
            if uid() and qcm.get("id"):
                sql("INSERT INTO tentatives(user_id,qcm_id,score,total,cree) VALUES(?,?,?,?,?)",
                    (uid(), qcm["id"], sum(r == q["r"] for r, q in zip(reps, qs)), len(qs), maintenant()))
            st.rerun()
        return

    # Correction
    reps = st.session_state.get("reponses", [None] * len(qs))
    score = sum(r == q["r"] for r, q in zip(reps, qs))
    note = round(score / len(qs) * 20, 1)
    m1, m2, m3 = st.columns(3)
    m1.metric("Ta note", f"{note} / 20")
    m2.metric("Bonnes réponses", f"{score} / {len(qs)}")
    m3.metric("Sans réponse", sum(r is None for r in reps))

    par_chap = {}
    for r, q in zip(reps, qs):
        a = par_chap.setdefault(q.get("chap", ""), [0, 0])
        a[0] += r == q["r"]
        a[1] += 1
    if len(par_chap) > 1:
        st.markdown("**Résultats par chapitre**")
        for ch, (b, t) in par_chap.items():
            st.progress(b / t, text=f"{ch} — {b}/{t}")
    st.divider()

    for i, (r, q) in enumerate(zip(reps, qs)):
        st.markdown(f"**{i + 1}.** {q['q']}")
        if r == q["r"]:
            st.success(f"✅ Correct : {LETTRES[r]}. {q['o'][r]}\n\n{q['e']}")
        else:
            ta = "⚠️ Sans réponse" if r is None else f"❌ Ta réponse : {LETTRES[r]}. {q['o'][r]}"
            st.error(f"{ta}\n\n**Bonne réponse : {LETTRES[q['r']]}. {q['o'][q['r']]}**\n\n{q['e']}")

    erreurs = [q for r, q in zip(reps, qs) if r != q["r"]]
    b1, b2 = st.columns(2)
    if b1.button("🔄 Recommencer", use_container_width=True):
        st.session_state.valide = False
        st.session_state.version += 1
        st.rerun()
    if erreurs and b2.button(f"🎯 Refaire mes {len(erreurs)} erreur(s)", use_container_width=True):
        st.session_state.qcm = {"id": None, "titre": "Révision de mes erreurs", "questions": erreurs}
        st.session_state.valide = False
        st.session_state.version += 1
        st.rerun()


SYSTEME_CHAT = """Tu es Étudiants Learn, un tuteur bienveillant, précis et pédagogue. L'étudiant t'a fourni son cours.
- Réponds en français, en t'appuyant d'abord sur le contenu du cours ; si tu sors du cours, dis-le clairement.
- Explique étape par étape, donne des exemples, et propose des exercices corrigés si on te le demande.
- Pour les maths, utilise LaTeX : $...$ en ligne et $$...$$ en bloc (jamais \\( \\) ni \\[ \\]).
- Sois concis quand la question est simple, détaillé quand elle est complexe."""


def page_chat(cle):
    entete("💬 Assistant", "Pose tes questions sur ton cours : explications, exercices, méthodes, révisions.")
    cours = besoin_cours()
    if not cours:
        return
    msgs = st.session_state.setdefault("chat", [])
    if not msgs:
        st.markdown("**Suggestions :**")
        sug = ["Explique-moi simplement les notions les plus importantes du cours.",
               "Donne-moi 3 exercices corrigés pas à pas sur ce cours.",
               "Quelles sont les erreurs classiques à éviter à l'examen ?"]
        cols = st.columns(len(sug))
        for col, s in zip(cols, sug):
            if col.button(s, use_container_width=True):
                st.session_state.question_en_attente = s
                st.rerun()
    for m in msgs:
        with st.chat_message("user" if m["role"] == "user" else "assistant"):
            st.markdown(m["content"])

    question = st.chat_input("Pose ta question sur le cours…") or st.session_state.pop("question_en_attente", None)
    if question:
        if not cle:
            st.error("Ajoute ta clé API Gemini.")
            return
        msgs.append({"role": "user", "content": question})
        with st.chat_message("user"):
            st.markdown(question)
        histo = msgs[-20:]
        contenu = [
            types.Content(role="user", parts=parts_cours(cours) + [
                types.Part.from_text(text="Voici mon cours. Sers-t'en comme base pour répondre à mes questions.")]),
            types.Content(role="model", parts=[types.Part.from_text(text="Compris, je suis prêt à t'aider.")]),
        ] + [types.Content(role="user" if m["role"] == "user" else "model",
                           parts=[types.Part.from_text(text=m["content"])]) for m in histo]
        with st.chat_message("assistant"):
            try:
                rep = st.write_stream(flux_chat(contenu, cle, SYSTEME_CHAT))
                msgs.append({"role": "assistant", "content": normaliser_math(rep)})
                sauver_chat(cours.get("id"), msgs)
                st.rerun()
            except Exception as e:
                msgs.pop()
                st.error(f"Erreur : {e}")
    if msgs and st.button("🧹 Effacer la conversation"):
        st.session_state.chat = []
        sauver_chat(cours.get("id"), [])
        st.rerun()


def page_historique():
    entete("🗂️ Historique", "Retrouve tes résumés, QCM et résultats.")
    if not uid():
        st.info("Crée un compte pour conserver ton historique.")
        return
    t1, t2 = st.tabs(["📝 Résumés", "✅ QCM et notes"])
    with t1:
        lst = sql("SELECT id, titre, cree FROM resumes WHERE user_id=? ORDER BY id DESC", (uid(),), fetch="all")
        if not lst:
            st.caption("Aucun résumé enregistré.")
        for r in lst:
            a, b, c = st.columns([5, 1.3, 1.3])
            a.markdown(f"**{r['titre']}**  \n<small>{r['cree']}</small>", unsafe_allow_html=True)
            if b.button("Ouvrir", key=f"hr{r['id']}", use_container_width=True):
                row = sql("SELECT contenu FROM resumes WHERE id=? AND user_id=?", (r["id"], uid()), fetch="one")
                d = json.loads(row["contenu"])
                d["id"] = r["id"]
                st.session_state.resume = d
                st.success("Résumé chargé : ouvre l'onglet **📝 Résumé**.")
            if c.button("🗑️", key=f"dr{r['id']}", use_container_width=True):
                sql("DELETE FROM resumes WHERE id=? AND user_id=?", (r["id"], uid()))
                st.rerun()
    with t2:
        lst = sql("""SELECT q.id, q.titre, q.cree, COUNT(t.id) AS essais,
                     MAX(CAST(t.score AS REAL) / t.total * 20) AS meilleure
                     FROM qcms q LEFT JOIN tentatives t ON t.qcm_id = q.id
                     WHERE q.user_id=? GROUP BY q.id ORDER BY q.id DESC""", (uid(),), fetch="all")
        if not lst:
            st.caption("Aucun QCM enregistré.")
        for q in lst:
            a, b, c = st.columns([5, 1.3, 1.3])
            meilleure = f"meilleure note {round(q['meilleure'], 1)}/20" if q["meilleure"] is not None else "pas encore tenté"
            a.markdown(f"**{q['titre']}**  \n<small>{q['cree']} · {q['essais']} essai(s) · {meilleure}</small>",
                       unsafe_allow_html=True)
            if b.button("Refaire", key=f"hq{q['id']}", use_container_width=True):
                row = sql("SELECT data FROM qcms WHERE id=? AND user_id=?", (q["id"], uid()), fetch="one")
                st.session_state.qcm = {"id": q["id"], "titre": q["titre"], "questions": json.loads(row["data"])}
                st.session_state.valide = False
                st.session_state.version = st.session_state.get("version", 0) + 1
                st.success("QCM chargé : ouvre l'onglet **✅ QCM**.")
            if c.button("🗑️", key=f"dq{q['id']}", use_container_width=True):
                sql("DELETE FROM tentatives WHERE qcm_id=? AND user_id=?", (q["id"], uid()))
                sql("DELETE FROM qcms WHERE id=? AND user_id=?", (q["id"], uid()))
                st.rerun()


# ----------------------------------------------------------------------------
# Application
# ----------------------------------------------------------------------------
init_db()

if "user" not in st.session_state:
    page_connexion()
    st.stop()

with st.sidebar:
    st.markdown('<div class="marque"><div class="logo">📚</div><b>Étudiants Learn</b></div>', unsafe_allow_html=True)
    st.caption(f"👤 {st.session_state.user['nom']}")
    page = st.radio("Navigation", PAGES, label_visibility="collapsed", key="nav")
    st.divider()
    if st.session_state.get("cours"):
        st.markdown("**Cours actif**")
        st.caption(st.session_state.cours["nom"])
    cle = cle_api()
    st.markdown(
        f'''<div class="promo">🎓 <b>Répétitions</b> : {TEL_REP[0]}<br>
        🗂️ <b>Démarches en ligne</b> (nationalité, passeport, casier judiciaire) &amp; traitement de données :
        {TEL_REP[0]} / {TEL_AMB[0]}</div>''', unsafe_allow_html=True)
    st.write("")
    if st.button("Se déconnecter", use_container_width=True):
        for k in list(st.session_state.keys()):
            del st.session_state[k]
        st.rerun()

if page == PAGES[0]:
    page_cours(cle)
elif page == PAGES[1]:
    page_resume(cle)
elif page == PAGES[2]:
    page_qcm(cle)
elif page == PAGES[3]:
    page_chat(cle)
elif page == PAGES[4]:
    page_historique()
else:
    page_services()

pied()
